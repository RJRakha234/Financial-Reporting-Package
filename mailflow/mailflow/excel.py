"""Spreadsheet-driven mail: each row of an ``.xlsx`` is one e-mail.

The spreadsheet is both the **input** (what to send, to whom, and when) and the
live **status board** (the tool writes Sent/Received status, the reply sender,
and the archive paths back into the same rows).

Recognised input columns (header names are case-insensitive; several aliases
are accepted):

    To            recipient address(es), separated by , or ;
    CC            optional cc address(es)
    Subject       the subject line (supports {Column} placeholders)
    Body          the message body  (supports {Column} placeholders)
    Attachments   file path(s)/glob(s), separated by ; or new lines
    Send Time     when to send (blank = as soon as the tool runs)

Any *other* column can be referenced as a ``{Column Name}`` placeholder in the
Subject or Body — i.e. mail-merge straight from the sheet.

Columns the tool fills in (created automatically if absent):

    Sent Status · Sent At · Message ID · Sent Path
    Received Status · Received From · Received At · Received Path · Notes

Server credentials never live in the spreadsheet — they come from the YAML
config / environment, exactly as for the rest of mailflow.
"""

from __future__ import annotations

import logging
import os
import re
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Config, Job, SendingPolicy
from .message import build_message, render, save_message, save_raw
from .replies import ImapReplyChecker
from .sender import Sender

log = logging.getLogger("mailflow")

# -- column resolution -------------------------------------------------------

INPUT_ALIASES: dict[str, list[str]] = {
    "to": ["to", "recipient", "recipients", "to recipients", "send to", "mail to"],
    "cc": ["cc", "cc recipients", "carbon copy"],
    "subject": ["subject", "mail subject", "title"],
    "body": ["body", "message", "mail body", "content", "mail content", "mail"],
    "attachments": ["attachments", "attachment", "files", "file", "attachment path"],
    "send_time": ["send time", "timer", "schedule", "send at", "scheduled time",
                  "send date", "time", "when"],
    "save_dir": ["save dir", "save folder", "sent folder", "sent dir"],
    "approved": ["approved?", "approved", "approval", "approve", "sign off"],
}

OUTPUT_HEADERS: dict[str, str] = {
    "sent_status": "Sent Status",
    "sent_at": "Sent At",
    "message_id": "Message ID",
    "sent_path": "Sent Path",
    "attempts": "Attempts",
    "received_status": "Received Status",
    "received_from": "Received From",
    "received_at": "Received At",
    "received_path": "Received Path",
    "reminders_sent": "Reminders Sent",
    "last_reminder": "Last Reminder",
    "error": "Notes",
}

_SENT = "Sent"
_FAILED = "Failed"
_AWAITING = "Awaiting"
_RECEIVED = "Received"
_HELD = "Held"  # awaiting approval

_TRUTHY = {"y", "yes", "true", "1", "approved", "ok", "✓", "done", "sign off", "signed"}

# Status-cell fills for the visual board.
_FILL = {
    _SENT: PatternFill("solid", fgColor="C6EFCE"),       # green
    _FAILED: PatternFill("solid", fgColor="FFC7CE"),     # red
    _AWAITING: PatternFill("solid", fgColor="FFEB9C"),   # amber
    _RECEIVED: PatternFill("solid", fgColor="BDD7EE"),   # blue
    _HELD: PatternFill("solid", fgColor="E2D9F3"),       # lilac
}


class ExcelError(Exception):
    """Raised for spreadsheet structure problems."""


@dataclass
class SendSummary:
    sent: int = 0
    failed: int = 0
    skipped_not_due: int = 0
    already_sent: int = 0
    held_for_approval: int = 0
    capped: int = 0


@dataclass
class ReplySummary:
    checked: int = 0
    received: int = 0


@dataclass
class ReminderSummary:
    reminded: int = 0
    escalated: int = 0


@dataclass
class StatusSummary:
    total: int = 0
    sent: int = 0
    awaiting: int = 0
    received: int = 0
    failed: int = 0
    held: int = 0
    overdue: int = 0


# -- spreadsheet wrapper -----------------------------------------------------


class Spreadsheet:
    """Thin openpyxl wrapper: header lookup, cell get/set, column creation."""

    def __init__(self, path: str | Path, sheet: str | None = None):
        self.path = Path(path)
        if not self.path.is_file():
            raise ExcelError(f"spreadsheet not found: {self.path}")
        self.wb = openpyxl.load_workbook(self.path)
        self.ws = self.wb[sheet] if sheet else self.wb.active
        self.headers: dict[str, int] = {}
        for col in range(1, self.ws.max_column + 1):
            val = self.ws.cell(row=1, column=col).value
            if val is not None and str(val).strip():
                self.headers[str(val).strip().lower()] = col

    def col(self, *names: str) -> int | None:
        for name in names:
            idx = self.headers.get(name.strip().lower())
            if idx:
                return idx
        return None

    def ensure_col(self, header_text: str) -> int:
        idx = self.col(header_text)
        if idx:
            return idx
        idx = self.ws.max_column + 1
        cell = self.ws.cell(row=1, column=idx, value=header_text)
        cell.font = Font(bold=True)
        self.headers[header_text.strip().lower()] = idx
        return idx

    def data_rows(self) -> list[int]:
        return list(range(2, self.ws.max_row + 1))

    def get(self, row: int, col: int | None) -> object:
        return self.ws.cell(row=row, column=col).value if col else None

    def set(self, row: int, col: int, value) -> None:
        self.ws.cell(row=row, column=col, value=value)

    def set_status(self, row: int, col: int, status: str) -> None:
        """Write a status string and colour the cell to match the board."""
        cell = self.ws.cell(row=row, column=col, value=status)
        fill = _FILL.get(status)
        if fill is not None:
            cell.fill = fill

    def row_dict(self) -> dict[int, str]:
        """Map column index -> original header text (for placeholders)."""
        out = {}
        for col in range(1, self.ws.max_column + 1):
            val = self.ws.cell(row=1, column=col).value
            if val is not None and str(val).strip():
                out[col] = str(val).strip()
        return out

    def save(self, path: str | Path | None = None) -> Path:
        target = Path(path) if path else self.path
        self.wb.save(target)
        return target


class FileLock:
    """A simple exclusive lock so two runs can't clobber the same workbook.

    Used as a context manager around an Excel operation. The lock is a sibling
    ``<file>.lock`` created atomically; a stale lock can be removed by hand.
    """

    def __init__(self, target: str | Path):
        self.lock_path = Path(str(target) + ".lock")
        self._fd: int | None = None

    def __enter__(self) -> "FileLock":
        try:
            self._fd = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(self._fd, str(os.getpid()).encode())
        except FileExistsError:
            raise ExcelError(
                f"{self.lock_path.name} exists — another mailflow run may be "
                f"using {self.lock_path.stem}. Remove the lock file if it is stale."
            ) from None
        return self

    def __exit__(self, *exc: object) -> None:
        if self._fd is not None:
            os.close(self._fd)
            self._fd = None
        self.lock_path.unlink(missing_ok=True)


# -- value parsing -----------------------------------------------------------


def split_addresses(value: object) -> list[str]:
    if value is None:
        return []
    parts = re.split(r"[,;\n]+", str(value))
    return [p.strip() for p in parts if p.strip()]


def split_paths(value: object) -> list[str]:
    if value is None:
        return []
    parts = re.split(r"[;\n]+", str(value))
    return [p.strip() for p in parts if p.strip()]


_TIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%d/%m/%Y %H:%M", "%d/%m/%Y", "%m/%d/%Y %H:%M", "%m/%d/%Y",
    "%d-%m-%Y %H:%M", "%d-%m-%Y",
)


def parse_when(value: object) -> datetime | None:
    """Parse a Send Time cell. Blank means 'send now'."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    s = str(value).strip()
    if s.lower() in {"now", "asap", "immediately"}:
        return None
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        raise ExcelError(
            f"unrecognised Send Time {value!r}; use e.g. '2026-06-22 08:00'"
        ) from None


def _context(sheet: Spreadsheet, row: int, now: datetime) -> dict[str, str]:
    """Build the {placeholder} context from every column of the row."""
    ctx: dict[str, str] = {
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M"),
        "datetime": now.strftime("%Y-%m-%d %H:%M"),
    }
    for col, header in sheet.row_dict().items():
        val = sheet.get(row, col)
        text = "" if val is None else str(val)
        # Expose several spellings so {First Name}/{First_Name}/{first_name} all work.
        for key in {header, header.replace(" ", "_"), header.lower(),
                    header.lower().replace(" ", "_")}:
            ctx[key] = text
    return ctx


# -- sending -----------------------------------------------------------------


def _require_columns(sheet: Spreadsheet) -> dict[str, int | None]:
    cols = {key: sheet.col(*aliases) for key, aliases in INPUT_ALIASES.items()}
    missing = [k for k in ("to", "subject") if cols[k] is None]
    if missing:
        found = ", ".join(sorted(sheet.headers)) or "(none)"
        raise ExcelError(
            f"spreadsheet is missing required column(s): {missing}. "
            f"Headers found: {found}"
        )
    return cols


def _truthy(value: object) -> bool:
    return value is not None and str(value).strip().lower() in _TRUTHY


def _send_with_retry(
    sender: Sender, msg, recipients: list[str], policy: SendingPolicy
) -> tuple[int, Exception | None]:
    """Try to send, retrying transient failures with linear backoff."""
    last: Exception | None = None
    for attempt in range(1, policy.max_attempts + 1):
        try:
            sender.send(msg, recipients)
            return attempt, None
        except Exception as exc:  # noqa: BLE001
            last = exc
            log.warning("send attempt %d/%d failed: %s",
                        attempt, policy.max_attempts, exc)
            if attempt < policy.max_attempts and policy.retry_backoff_seconds > 0:
                time.sleep(policy.retry_backoff_seconds * attempt)
    return policy.max_attempts, last


def send_due(
    sheet: Spreadsheet,
    config: Config,
    sender: Sender,
    now: datetime | None = None,
    dry_run: bool = False,
) -> SendSummary:
    """Send every approved, due row that has not been sent yet.

    Honours the approval gate, per-run rate cap, retry policy, and throttle
    from ``config.sending``, and colour-codes the status cells.
    """
    now = now or datetime.now()
    cols = _require_columns(sheet)
    out = {k: sheet.ensure_col(h) for k, h in OUTPUT_HEADERS.items()} if not dry_run else {}
    policy = config.sending
    summary = SendSummary()
    sent_this_run = 0

    for row in sheet.data_rows():
        to = split_addresses(sheet.get(row, cols["to"]))
        subject = sheet.get(row, cols["subject"])
        if not to or not subject:
            continue  # blank/spacer row

        status_col = out.get("sent_status") or sheet.col(OUTPUT_HEADERS["sent_status"])
        status = str(sheet.get(row, status_col) or "").strip().lower()
        if status == _SENT.lower():
            summary.already_sent += 1
            continue

        when = parse_when(sheet.get(row, cols["send_time"])) if cols["send_time"] else None
        if when is not None and when > now:
            summary.skipped_not_due += 1
            continue

        # Approval gate: if an approval column exists, the row must be approved.
        if cols["approved"] is not None and not _truthy(sheet.get(row, cols["approved"])):
            summary.held_for_approval += 1
            if not dry_run:
                sheet.set_status(row, out["sent_status"], _HELD)
                sheet.set(row, out["error"], "awaiting approval")
            continue

        # Per-run rate cap (protects against corporate send limits).
        if policy.max_per_run is not None and sent_this_run >= policy.max_per_run:
            summary.capped += 1
            continue

        job = Job(
            name=f"row{row}",
            to=to,
            cc=split_addresses(sheet.get(row, cols["cc"])) if cols["cc"] else [],
            subject=str(subject),
            body=str(sheet.get(row, cols["body"]) or "") if cols["body"] else "",
            attachments=split_paths(sheet.get(row, cols["attachments"])) if cols["attachments"] else [],
            variables=_context(sheet, row, now),
        )

        if dry_run:
            log.info("[dry-run] row %s -> %s | %s", row, ", ".join(to), job.subject)
            summary.sent += 1
            sent_this_run += 1
            continue

        # Build + archive first; a bad attachment fails the row cleanly.
        try:
            msg = build_message(job, config.smtp, now)
            override = sheet.get(row, cols["save_dir"]) if cols["save_dir"] else None
            save_dir = str(override) if override else config.save_dir
            saved = save_message(msg, save_dir, f"row{row}")
        except Exception as exc:  # noqa: BLE001
            sheet.set_status(row, out["sent_status"], _FAILED)
            sheet.set(row, out["error"], str(exc)[:300])
            summary.failed += 1
            log.error("row %s could not be prepared: %s", row, exc)
            continue

        attempts, error = _send_with_retry(sender, msg, job.all_recipients(), policy)
        sheet.set(row, out["attempts"], attempts)
        if error is not None:
            sheet.set_status(row, out["sent_status"], _FAILED)
            sheet.set(row, out["error"], str(error)[:300])
            summary.failed += 1
            log.error("row %s failed after %d attempt(s): %s", row, attempts, error)
            continue

        sheet.set_status(row, out["sent_status"], _SENT)
        sheet.set(row, out["sent_at"], now.strftime("%Y-%m-%d %H:%M:%S"))
        sheet.set(row, out["message_id"], msg["Message-ID"])
        sheet.set(row, out["sent_path"], str(saved))
        sheet.set_status(row, out["received_status"], _AWAITING)
        sheet.set(row, out["error"], "")
        summary.sent += 1
        sent_this_run += 1
        log.info("row %s sent to %s (saved %s)", row, ", ".join(to), saved)
        if policy.throttle_seconds > 0:
            time.sleep(policy.throttle_seconds)

    if not dry_run:
        sheet.save()
    return summary


# -- revert detection --------------------------------------------------------


def check_replies(
    sheet: Spreadsheet,
    config: Config,
    checker: ImapReplyChecker | None = None,
    now: datetime | None = None,
) -> ReplySummary:
    """Look for replies to sent rows; archive them and write status back."""
    if config.imap is None:
        raise ExcelError("no 'imap' section in config; cannot detect reverts")
    now = now or datetime.now()
    checker = checker or ImapReplyChecker(config.imap)

    out = {k: sheet.ensure_col(h) for k, h in OUTPUT_HEADERS.items()}
    summary = ReplySummary()

    conn = checker.connect()
    try:
        for row in sheet.data_rows():
            sent = str(sheet.get(row, out["sent_status"]) or "").strip().lower()
            recvd = str(sheet.get(row, out["received_status"]) or "").strip().lower()
            msg_id = str(sheet.get(row, out["message_id"]) or "").strip()
            if sent != _SENT.lower() or recvd == _RECEIVED.lower() or not msg_id:
                continue
            summary.checked += 1
            hit = checker.find_reply(conn, msg_id)
            if hit is None:
                continue
            saved = save_raw(hit.raw, config.received_dir, f"row{row}-revert")
            sheet.set_status(row, out["received_status"], _RECEIVED)
            sheet.set(row, out["received_from"], hit.from_addr)
            sheet.set(row, out["received_at"], hit.received_at or now.strftime("%Y-%m-%d %H:%M:%S"))
            sheet.set(row, out["received_path"], str(saved))
            summary.received += 1
            log.info("row %s revert from %s (saved %s)", row, hit.from_addr, saved)
    finally:
        checker._logout(conn)

    sheet.save()
    return summary


# -- reminders / SLA escalation ---------------------------------------------


def send_reminders(
    sheet: Spreadsheet,
    config: Config,
    sender: Sender,
    now: datetime | None = None,
) -> ReminderSummary:
    """Chase rows that have been Awaiting a revert beyond the configured SLA.

    Sends up to ``max_reminders`` reminders per row (no more often than
    ``every_hours``), cc-ing ``escalate_to`` on the final reminder.
    """
    policy = config.reminders
    summary = ReminderSummary()
    if not policy.enabled:
        return summary

    now = now or datetime.now()
    cols = _require_columns(sheet)
    out = {k: sheet.ensure_col(h) for k, h in OUTPUT_HEADERS.items()}
    sla = timedelta(hours=policy.sla_hours)
    gap = timedelta(hours=policy.every_hours)

    for row in sheet.data_rows():
        sent = str(sheet.get(row, out["sent_status"]) or "").strip().lower()
        recvd = str(sheet.get(row, out["received_status"]) or "").strip().lower()
        if sent != _SENT.lower() or recvd == _RECEIVED.lower():
            continue

        sent_at = parse_when(sheet.get(row, out["sent_at"]))
        if sent_at is None or now - sent_at < sla:
            continue  # not overdue yet

        reminders_sent = int(sheet.get(row, out["reminders_sent"]) or 0)
        if reminders_sent >= policy.max_reminders:
            continue  # exhausted

        last_reminder = parse_when(sheet.get(row, out["last_reminder"]))
        if last_reminder is not None and now - last_reminder < gap:
            continue  # too soon since the last nudge

        to = split_addresses(sheet.get(row, cols["to"]))
        if not to:
            continue
        is_final = reminders_sent + 1 >= policy.max_reminders
        cc = list(policy.escalate_to) if (is_final and policy.escalate_to) else []

        context = _context(sheet, row, now)
        reminder = Job(
            name=f"row{row}-reminder",
            to=to,
            cc=cc,
            subject=render(policy.subject, context),
            body=render(policy.body, context),
            variables=context,
        )
        try:
            msg = build_message(reminder, config.smtp, now)
            save_message(msg, config.save_dir, f"row{row}-reminder")
            attempts, error = _send_with_retry(
                sender, msg, reminder.all_recipients(), config.sending
            )
            if error is not None:
                raise error
        except Exception as exc:  # noqa: BLE001
            log.error("row %s reminder failed: %s", row, exc)
            continue

        sheet.set(row, out["reminders_sent"], reminders_sent + 1)
        sheet.set(row, out["last_reminder"], now.strftime("%Y-%m-%d %H:%M:%S"))
        summary.reminded += 1
        if cc:
            summary.escalated += 1
            log.info("row %s escalated to %s", row, ", ".join(cc))
        else:
            log.info("row %s reminded (%d)", row, reminders_sent + 1)

    sheet.save()
    return summary


# -- status board ------------------------------------------------------------


def summarize(sheet: Spreadsheet, config: Config, now: datetime | None = None) -> StatusSummary:
    """Tally the board: counts per status and how many are overdue for a revert."""
    now = now or datetime.now()
    s_col = sheet.col(OUTPUT_HEADERS["sent_status"])
    r_col = sheet.col(OUTPUT_HEADERS["received_status"])
    at_col = sheet.col(OUTPUT_HEADERS["sent_at"])
    to_col = sheet.col(*INPUT_ALIASES["to"])
    sla = timedelta(hours=config.reminders.sla_hours)
    summary = StatusSummary()

    for row in sheet.data_rows():
        if not (to_col and sheet.get(row, to_col)):
            continue
        summary.total += 1
        sent = str(sheet.get(row, s_col) or "").strip().lower() if s_col else ""
        recvd = str(sheet.get(row, r_col) or "").strip().lower() if r_col else ""
        if sent == _SENT.lower():
            summary.sent += 1
        elif sent == _FAILED.lower():
            summary.failed += 1
        elif sent == _HELD.lower():
            summary.held += 1
        if recvd == _RECEIVED.lower():
            summary.received += 1
        elif sent == _SENT.lower():
            summary.awaiting += 1
            sent_at = parse_when(sheet.get(row, at_col)) if at_col else None
            if sent_at is not None and now - sent_at >= sla:
                summary.overdue += 1
    return summary


# -- template ----------------------------------------------------------------


def make_template(path: str | Path) -> Path:
    """Write a ready-to-edit .xlsx with headers and one example row."""
    path = Path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mails"
    headers = [
        "To", "CC", "Subject", "Body", "Attachments", "Send Time", "Approved?",
        *OUTPUT_HEADERS.values(),
    ]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)
    ws.append([
        "cfo@example.com, controller@example.com",
        "audit@example.com",
        "Weekly Financial Report - {date}",
        "Hi team,\n\nPlease find this week's report attached. Kindly revert with approval.",
        "reports/weekly_latest.pdf",
        "2026-06-22 08:00",
        "No",
    ])
    wide = {"To": 34, "CC": 22, "Subject": 30, "Body": 40, "Attachments": 26,
            "Send Time": 18, "Sent Path": 26, "Received Path": 26, "Notes": 28}
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = wide.get(header, 15)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path
