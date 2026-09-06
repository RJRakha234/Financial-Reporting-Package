"""Live demo of the Excel-driven mode with the efficiency features.

Shows, end to end, against the local capture relay:
  * one row per mail, with a mail-merge {Region} placeholder;
  * an approval gate (an un-approved row is HELD, not sent);
  * status written back + colour-coded into the sheet;
  * revert detection (simulated inbound reply);
  * SLA reminders with escalation for a row still awaiting a revert;
  * a status-board summary.

The IMAP matching logic is real and unit-tested; here it is fed one canned
reply so no live mailbox is needed.

    python excel_demo.py --port 8025
"""

from __future__ import annotations

import argparse
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from mailflow import Spreadsheet, send_due, send_reminders, summarize
from mailflow.config import (
    Config,
    ImapConfig,
    ReminderPolicy,
    SendingPolicy,
    SmtpConfig,
)
from mailflow.excel import check_replies
from mailflow.replies import ReplyHit
from mailflow.sender import SmtpSender

SHEET = "acme_mails.xlsx"
SHOW = ["To", "Subject", "Approved?", "Sent Status", "Received Status",
        "Received From", "Reminders Sent"]


def build_sheet() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mails"
    headers = ["To", "CC", "Subject", "Body", "Send Time", "Approved?", "Region"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).font = Font(bold=True)
    ws.append(["cfo@acme.example, controller@acme.example", "audit@acme.example",
               "Acme Weekly Report - {Region} - {date}",
               "Hi team, please find the {Region} figures attached. Kindly revert.",
               "", "Yes", "Group"])
    ws.append(["ap@acme.example", "",
               "Vendor Statement - {Region}",
               "Please reconcile the attached {Region} vendor statement.",
               "", "Yes", "EMEA"])
    ws.append(["board@acme.example", "",
               "Board Pack - {Region}",
               "Confidential board pack for review.",
               "", "No", "Group"])   # not approved -> Held
    wb.save(SHEET)


def dump(title: str) -> None:
    wb = openpyxl.load_workbook(SHEET)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    w = {"To": 34, "Subject": 30, "Approved?": 10, "Sent Status": 12,
         "Received Status": 16, "Received From": 18, "Reminders Sent": 14}
    print(f"\n--- spreadsheet: {title} ---")
    print(" | ".join(n[: w[n]].ljust(w[n]) for n in SHOW))
    print("-" * 150)
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, head["To"]).value:
            continue
        cells = []
        for n in SHOW:
            v = ws.cell(r, head[n]).value if n in head else ""
            cells.append(("" if v is None else str(v))[: w[n]].ljust(w[n]))
        print(" | ".join(cells))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--port", type=int, default=8025)
    args = ap.parse_args()

    config = Config(
        smtp=SmtpConfig(host="127.0.0.1", port=args.port, security="none", auth=False,
                        username="finance-automation@acme.example", password_env="",
                        from_addr="finance-automation@acme.example"),
        imap=ImapConfig(host="imap.local", port=993, security="ssl",
                        username="x", password_env="X"),
        jobs=[], save_dir="archive_sent", received_dir="archive_received",
        sending=SendingPolicy(max_attempts=3, retry_backoff_seconds=0),
        reminders=ReminderPolicy(enabled=True, sla_hours=48, max_reminders=2,
                                 every_hours=0, escalate_to=["cfo-office@acme.example"]),
    )

    print("=== 1. Build the control spreadsheet (one row = one mail) ===")
    build_sheet()
    dump("BEFORE — three mails, board pack NOT approved")

    print("\n=== 2. Send approved, due rows (board pack is held back) ===")
    sheet = Spreadsheet(SHEET)
    s = send_due(sheet, config, SmtpSender(config.smtp))
    print(f"send summary: sent={s.sent} held-for-approval={s.held_for_approval}")
    dump("AFTER SEND — 2 Sent (Awaiting), board pack Held")

    print("\n=== 3. The CFO reverts. Reconcile reverts from the mailbox ===")
    wb = openpyxl.load_workbook(SHEET)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    cfo_msgid = ws.cell(2, head["Message ID"]).value

    class FakeChecker:
        def connect(self): return object()
        def _logout(self, conn): pass
        def find_reply(self, conn, message_id):
            if message_id == cfo_msgid:
                return ReplyHit(from_addr="cfo@acme.example",
                                raw=b"From: CFO <cfo@acme.example>\r\n"
                                    b"Subject: Re: Acme Weekly Report\r\n\r\nApproved.",
                                received_at="2026-06-19 09:41:00")
            return None

    r = check_replies(Spreadsheet(SHEET), config, checker=FakeChecker())
    print(f"revert summary: received={r.received}")
    dump("AFTER REVERT — CFO row Received; AP still Awaiting")

    print("\n=== 4. AP has not reverted in 72h (> 48h SLA). Send a reminder ===")
    # Back-date the AP row's Sent At so it is overdue.
    wb = openpyxl.load_workbook(SHEET)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    overdue = (datetime.now() - timedelta(hours=72)).strftime("%Y-%m-%d %H:%M:%S")
    ws.cell(3, head["Sent At"]).value = overdue
    wb.save(SHEET)

    rem = send_reminders(Spreadsheet(SHEET), config, SmtpSender(config.smtp))
    print(f"reminder summary: reminded={rem.reminded} escalated={rem.escalated}")
    dump("AFTER REMINDER — AP row 'Reminders Sent' incremented")

    print("\n=== 5. Status board summary ===")
    st = summarize(Spreadsheet(SHEET), config)
    print(f"  total={st.total}  sent={st.sent}  received={st.received}  "
          f"awaiting={st.awaiting} (overdue={st.overdue})  held={st.held}  failed={st.failed}")

    print("\n=== 6. On-disk archives ===")
    for label, d in [("sent", "archive_sent"), ("received", "archive_received")]:
        files = sorted(Path(d).glob("*.eml")) if Path(d).exists() else []
        print(f"  {label}: " + (", ".join(f.name for f in files) or "(none)"))


if __name__ == "__main__":
    main()
