from datetime import datetime, timedelta

import openpyxl
import pytest

from mailflow import (
    DryRunSender,
    Spreadsheet,
    send_due,
    send_reminders,
    summarize,
)
from mailflow.config import Config, ImapConfig, ReminderPolicy, SendingPolicy, SmtpConfig
from mailflow.excel import FileLock


def make_config(tmp_path, sending=None, reminders=None, imap=False):
    return Config(
        smtp=SmtpConfig(host="h", port=25, security="none", username="bot@corp.com",
                        password_env="X", auth=False),
        imap=ImapConfig(host="i", port=993, security="ssl", username="u",
                        password_env="Y") if imap else None,
        jobs=[],
        save_dir=str(tmp_path / "sent"),
        received_dir=str(tmp_path / "recv"),
        sending=sending or SendingPolicy(max_attempts=3, retry_backoff_seconds=0),
        reminders=reminders or ReminderPolicy(),
    )


def make_sheet(tmp_path, headers, rows, name="m.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    path = tmp_path / name
    wb.save(path)
    return path


def reload(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    return ws, head


# -- approval gate -----------------------------------------------------------


def test_approval_gate_holds_unapproved(tmp_path):
    cfg = make_config(tmp_path)
    path = make_sheet(
        tmp_path,
        ["To", "Subject", "Body", "Send Time", "Approved?"],
        [
            {"To": "a@e.com", "Subject": "Go", "Body": "x", "Approved?": "Yes"},
            {"To": "b@e.com", "Subject": "Wait", "Body": "x", "Approved?": "No"},
            {"To": "c@e.com", "Subject": "Blank", "Body": "x", "Approved?": ""},
        ],
    )
    sender = DryRunSender()
    s = send_due(Spreadsheet(path), cfg, sender)
    assert s.sent == 1 and s.held_for_approval == 2
    assert sender.sent[0][0] == "Go"

    ws, head = reload(path)
    assert ws.cell(2, head["Sent Status"]).value == "Sent"
    assert ws.cell(3, head["Sent Status"]).value == "Held"
    assert ws.cell(4, head["Sent Status"]).value == "Held"

    # Approving the held row and re-running sends it.
    ws.cell(3, head["Approved?"]).value = "yes"
    wb = ws.parent
    wb.save(path)
    s2 = send_due(Spreadsheet(path), cfg, DryRunSender())
    assert s2.sent == 1 and s2.already_sent == 1


# -- retry -------------------------------------------------------------------


class FlakySender:
    """Fails the first ``fail_times`` calls, then succeeds."""

    def __init__(self, fail_times):
        self.fail_times = fail_times
        self.calls = 0

    def send(self, msg, recipients):
        self.calls += 1
        if self.calls <= self.fail_times:
            raise RuntimeError("temporary glitch")


def test_retry_succeeds_then_records_attempts(tmp_path):
    cfg = make_config(tmp_path, sending=SendingPolicy(max_attempts=3, retry_backoff_seconds=0))
    path = make_sheet(tmp_path, ["To", "Subject", "Body"],
                      [{"To": "a@e.com", "Subject": "s", "Body": "b"}])
    s = send_due(Spreadsheet(path), cfg, FlakySender(fail_times=2))
    assert s.sent == 1 and s.failed == 0
    ws, head = reload(path)
    assert ws.cell(2, head["Attempts"]).value == 3
    assert ws.cell(2, head["Sent Status"]).value == "Sent"


def test_retry_exhausted_marks_failed(tmp_path):
    cfg = make_config(tmp_path, sending=SendingPolicy(max_attempts=2, retry_backoff_seconds=0))
    path = make_sheet(tmp_path, ["To", "Subject", "Body"],
                      [{"To": "a@e.com", "Subject": "s", "Body": "b"}])
    s = send_due(Spreadsheet(path), cfg, FlakySender(fail_times=5))
    assert s.sent == 0 and s.failed == 1
    ws, head = reload(path)
    assert ws.cell(2, head["Sent Status"]).value == "Failed"
    assert "glitch" in (ws.cell(2, head["Notes"]).value or "")


# -- rate cap ----------------------------------------------------------------


def test_max_per_run_caps_sending(tmp_path):
    cfg = make_config(tmp_path, sending=SendingPolicy(max_per_run=1, retry_backoff_seconds=0))
    path = make_sheet(tmp_path, ["To", "Subject", "Body"],
                      [{"To": f"{n}@e.com", "Subject": "s", "Body": "b"} for n in range(3)])
    s = send_due(Spreadsheet(path), cfg, DryRunSender())
    assert s.sent == 1 and s.capped == 2


# -- reminders ---------------------------------------------------------------


def _sent_sheet(tmp_path):
    """A sheet with one row already sent (Awaiting) and an old Sent At."""
    cfg = make_config(tmp_path)
    path = make_sheet(tmp_path, ["To", "Subject", "Body"],
                      [{"To": "client@e.com", "Subject": "Report", "Body": "b"}])
    send_due(Spreadsheet(path), cfg, DryRunSender())
    # Back-date the Sent At so it is past the SLA.
    ws, head = reload(path)
    old = (datetime.now() - timedelta(hours=72)).strftime("%Y-%m-%d %H:%M:%S")
    ws.cell(2, head["Sent At"]).value = old
    ws.parent.save(path)
    return path


def test_reminders_disabled_by_default(tmp_path):
    cfg = make_config(tmp_path)  # reminders.enabled is False
    path = _sent_sheet(tmp_path)
    r = send_reminders(Spreadsheet(path), cfg, DryRunSender())
    assert r.reminded == 0


def test_reminder_sent_when_overdue_and_escalates(tmp_path):
    policy = ReminderPolicy(enabled=True, sla_hours=48, max_reminders=1,
                            every_hours=0, escalate_to=["boss@e.com"])
    cfg = make_config(tmp_path, reminders=policy)
    path = _sent_sheet(tmp_path)
    sender = DryRunSender()
    r = send_reminders(Spreadsheet(path), cfg, sender)
    assert r.reminded == 1 and r.escalated == 1  # final reminder escalates
    # The boss was cc'd.
    assert "boss@e.com" in sender.sent[0][1]
    ws, head = reload(path)
    assert ws.cell(2, head["Reminders Sent"]).value == 1

    # Second run: max_reminders reached, nothing more.
    r2 = send_reminders(Spreadsheet(path), cfg, DryRunSender())
    assert r2.reminded == 0


# -- status board ------------------------------------------------------------


def test_summarize_counts(tmp_path):
    policy = ReminderPolicy(enabled=True, sla_hours=48, every_hours=0)
    cfg = make_config(tmp_path, reminders=policy)
    path = _sent_sheet(tmp_path)  # one sent+awaiting+overdue row
    s = summarize(Spreadsheet(path), cfg)
    assert s.total == 1 and s.sent == 1 and s.awaiting == 1 and s.overdue == 1


# -- file lock ---------------------------------------------------------------


def test_file_lock_is_exclusive(tmp_path):
    target = tmp_path / "m.xlsx"
    target.write_bytes(b"x")
    with FileLock(target):
        with pytest.raises(Exception):
            with FileLock(target):
                pass
    # Lock released on exit -> can acquire again.
    with FileLock(target):
        pass
