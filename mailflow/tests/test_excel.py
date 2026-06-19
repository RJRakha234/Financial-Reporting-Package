from datetime import datetime

import openpyxl
import pytest

from mailflow import DryRunSender, Spreadsheet, make_template, send_due
from mailflow.config import Config, ImapConfig, SmtpConfig
from mailflow.excel import ExcelError, check_replies, parse_when, split_addresses
from mailflow.replies import ReplyHit


def make_config(tmp_path, with_imap=False):
    smtp = SmtpConfig(
        host="h", port=25, security="none", username="bot@corp.com",
        password_env="X", auth=False,
    )
    imap = None
    if with_imap:
        imap = ImapConfig(
            host="i", port=993, security="ssl", username="bot@corp.com",
            password_env="Y",
        )
    return Config(
        smtp=smtp, imap=imap, jobs=[],
        save_dir=str(tmp_path / "sent"),
        received_dir=str(tmp_path / "recv"),
    )


def make_sheet(tmp_path, rows):
    """rows: list of dicts with keys To, CC, Subject, Body, Send Time."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["To", "CC", "Subject", "Body", "Send Time", "Region"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    path = tmp_path / "mails.xlsx"
    wb.save(path)
    return path


def test_helpers():
    assert split_addresses("a@e.com, b@e.com; c@e.com") == [
        "a@e.com", "b@e.com", "c@e.com"
    ]
    assert parse_when("") is None
    assert parse_when("2026-06-22 08:00") == datetime(2026, 6, 22, 8, 0)
    assert parse_when(datetime(2026, 1, 1, 9, 0)) == datetime(2026, 1, 1, 9, 0)
    with pytest.raises(ExcelError):
        parse_when("not a date")


def test_send_due_writes_status_back(tmp_path):
    cfg = make_config(tmp_path)
    path = make_sheet(tmp_path, [
        # due now (blank time), with a mail-merge placeholder {Region}
        {"To": "a@e.com", "CC": "c@e.com", "Subject": "Hello {Region}",
         "Body": "Hi {Region}", "Send Time": "", "Region": "APAC"},
        # scheduled far in the future -> not due
        {"To": "b@e.com", "Subject": "Later", "Body": "x",
         "Send Time": "2999-01-01 00:00", "Region": "EU"},
    ])
    sheet = Spreadsheet(path)
    sender = DryRunSender()
    summary = send_due(sheet, cfg, sender)

    assert summary.sent == 1
    assert summary.skipped_not_due == 1
    # The merge placeholder was resolved from the Region column.
    assert sender.sent[0][0] == "Hello APAC"
    assert sender.sent[0][1] == ["a@e.com", "c@e.com"]

    # Re-open the saved workbook and confirm the write-back.
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(2, head["Sent Status"]).value == "Sent"
    assert ws.cell(2, head["Received Status"]).value == "Awaiting"
    assert ws.cell(2, head["Message ID"]).value.endswith("@corp.com>")
    assert ws.cell(2, head["Sent Path"]).value.endswith(".eml")
    # Future row untouched.
    assert ws.cell(3, head["Sent Status"]).value in (None, "")


def test_send_due_is_idempotent(tmp_path):
    cfg = make_config(tmp_path)
    path = make_sheet(tmp_path, [
        {"To": "a@e.com", "Subject": "Once", "Body": "b", "Send Time": ""},
    ])
    sheet = Spreadsheet(path)
    assert send_due(sheet, cfg, DryRunSender()).sent == 1
    # Second run: already sent, nothing re-sent.
    sheet2 = Spreadsheet(path)
    s = send_due(sheet2, cfg, DryRunSender())
    assert s.sent == 0 and s.already_sent == 1


def test_missing_required_column_raises(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Recipient", "Body"])  # no Subject
    ws.append(["a@e.com", "x"])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(ExcelError):
        send_due(Spreadsheet(path), make_config(tmp_path), DryRunSender())


def test_dry_run_writes_nothing(tmp_path):
    cfg = make_config(tmp_path)
    path = make_sheet(tmp_path, [
        {"To": "a@e.com", "Subject": "s", "Body": "b", "Send Time": ""},
    ])
    sheet = Spreadsheet(path)
    send_due(sheet, cfg, DryRunSender(), dry_run=True)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # No "Sent Status" column was added on a dry run.
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert "Sent Status" not in headers


def test_template_roundtrips(tmp_path):
    path = make_template(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(path)
    headers = [c.value for c in wb.active[1]]
    for required in ("To", "Subject", "Sent Status", "Received Status"):
        assert required in headers


class _FakeChecker:
    """Stand-in for ImapReplyChecker that returns a canned reply."""

    def __init__(self, hit):
        self.hit = hit

    def connect(self):
        return object()

    def find_reply(self, conn, message_id):
        return self.hit

    def _logout(self, conn):
        pass


def test_check_replies_records_and_archives(tmp_path):
    cfg = make_config(tmp_path, with_imap=True)
    path = make_sheet(tmp_path, [
        {"To": "a@e.com", "Subject": "s", "Body": "b", "Send Time": ""},
    ])
    # Send first so the row has a Message ID + Awaiting status.
    send_due(Spreadsheet(path), cfg, DryRunSender())

    hit = ReplyHit(from_addr="a@e.com", raw=b"From: a@e.com\r\nSubject: Re: s\r\n\r\nok",
                   received_at="2026-06-19T10:00:00")
    sheet = Spreadsheet(path)
    summary = check_replies(sheet, cfg, checker=_FakeChecker(hit))
    assert summary.received == 1

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(2, head["Received Status"]).value == "Received"
    assert ws.cell(2, head["Received From"]).value == "a@e.com"
    assert ws.cell(2, head["Received Path"]).value.endswith(".eml")
