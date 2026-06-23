"""Unit tests for the XBRL instance + schema parsing and the workbook build.

The tests use tiny inline XBRL/XSD fixtures (no network, no large files) so they
exercise the join logic — facts -> contexts/units, and concepts -> labels.
"""

import textwrap
from pathlib import Path

import pytest

from xbrlextract import extract, parse_instance, parse_schema

INSTANCE = textwrap.dedent(
    """\
    <?xml version="1.0" encoding="utf-8"?>
    <xbrl xmlns="http://www.xbrl.org/2003/instance"
          xmlns:xbrldi="http://xbrl.org/2006/xbrldi"
          xmlns:iso4217="http://www.xbrl.org/2003/iso4217"
          xmlns:dei="http://xbrl.sec.gov/dei/2025"
          xmlns:ifrs-full="https://xbrl.ifrs.org/taxonomy/2025-03-27/ifrs-full"
          xmlns:infy="http://www.infosys.com/20260331">
      <context id="dur">
        <entity><identifier scheme="http://www.sec.gov/CIK">0001067491</identifier></entity>
        <period><startDate>2025-04-01</startDate><endDate>2026-03-31</endDate></period>
      </context>
      <context id="inst">
        <entity><identifier scheme="http://www.sec.gov/CIK">0001067491</identifier></entity>
        <period><instant>2026-03-31</instant></period>
      </context>
      <context id="seg">
        <entity><identifier scheme="http://www.sec.gov/CIK">0001067491</identifier>
          <segment><xbrldi:explicitMember dimension="ifrs-full:OperatingSegmentsAxis">infy:NorthAmericaMember</xbrldi:explicitMember></segment>
        </entity>
        <period><startDate>2025-04-01</startDate><endDate>2026-03-31</endDate></period>
      </context>
      <unit id="inr"><measure>iso4217:INR</measure></unit>
      <ifrs-full:Revenue contextRef="dur" unitRef="inr" decimals="-6" id="f1">1750000000000</ifrs-full:Revenue>
      <ifrs-full:Assets contextRef="inst" unitRef="inr" decimals="-6" id="f2">1480000000000</ifrs-full:Assets>
      <infy:SegmentRevenue contextRef="seg" unitRef="inr" decimals="-6" id="f3">1050000000000</infy:SegmentRevenue>
      <dei:DocumentType contextRef="dur" id="f4">20-F</dei:DocumentType>
    </xbrl>
    """
)

SCHEMA = textwrap.dedent(
    """\
    <?xml version="1.0" encoding="utf-8"?>
    <xsd:schema xmlns:xsd="http://www.w3.org/2001/XMLSchema"
                xmlns:link="http://www.xbrl.org/2003/linkbase"
                xmlns:xlink="http://www.w3.org/1999/xlink"
                xmlns:xbrli="http://www.xbrl.org/2003/instance"
                xmlns:infy="http://www.infosys.com/20260331"
                targetNamespace="http://www.infosys.com/20260331">
      <xsd:element name="SegmentRevenue" id="infy_SegmentRevenue"
                   type="xbrli:monetaryItemType" xbrli:periodType="duration"
                   xbrli:balance="credit" abstract="false"/>
      <xsd:annotation><xsd:appinfo>
        <link:labelLink xlink:type="extended"
                        xlink:role="http://www.xbrl.org/2003/role/link">
          <link:loc xlink:type="locator"
                    xlink:href="infy-20260331.xsd#infy_SegmentRevenue"
                    xlink:label="loc_seg"/>
          <link:label xlink:type="resource" xlink:label="lab_seg"
                      xlink:role="http://www.xbrl.org/2003/role/label">Segment revenue</link:label>
          <link:labelArc xlink:type="arc"
                         xlink:arcrole="http://www.xbrl.org/2003/arcrole/concept-label"
                         xlink:from="loc_seg" xlink:to="lab_seg"/>
        </link:labelLink>
      </xsd:appinfo></xsd:annotation>
    </xsd:schema>
    """
)


@pytest.fixture()
def files(tmp_path):
    inst = tmp_path / "inst_htm.xml"
    xsd = tmp_path / "inst.xsd"
    inst.write_text(INSTANCE)
    xsd.write_text(SCHEMA)
    return inst, xsd


def test_instance_parses_facts_contexts_units(files):
    inst = parse_instance(str(files[0]))
    assert len(inst.facts) == 4
    assert len(inst.contexts) == 3
    assert "inr" in inst.units
    assert inst.units["inr"].measure == "INR"


def test_fact_joins_to_context_and_unit(files):
    inst = parse_instance(str(files[0]))
    rev = next(f for f in inst.facts if f.localname == "Revenue")
    assert rev.concept == "ifrs-full:Revenue"
    ctx = inst.context_of(rev)
    assert ctx.period_type == "duration"
    assert ctx.start == "2025-04-01" and ctx.end == "2026-03-31"
    assert inst.unit_of(rev).measure == "INR"


def test_dimensional_context_is_resolved(files):
    inst = parse_instance(str(files[0]))
    seg = next(f for f in inst.facts if f.localname == "SegmentRevenue")
    dims = inst.context_of(seg).dimensions
    assert len(dims) == 1
    assert dims[0].axis == "ifrs-full:OperatingSegmentsAxis"
    assert dims[0].member == "infy:NorthAmericaMember"


def test_instant_vs_duration(files):
    inst = parse_instance(str(files[0]))
    assets = next(f for f in inst.facts if f.localname == "Assets")
    assert inst.context_of(assets).period_type == "instant"
    assert inst.context_of(assets).instant == "2026-03-31"


def test_schema_labels_and_concept_decl(files):
    tax = parse_schema(str(files[1]))
    assert tax.label("infy:SegmentRevenue") == "Segment revenue"
    c = tax.concepts["infy:SegmentRevenue"]
    assert c.extension and c.type == "monetaryItemType"
    assert c.period_type == "duration" and c.balance == "credit"


def test_label_fallback_humanizes_unknown_concept(files):
    tax = parse_schema(str(files[1]))
    # not in the (tiny) label linkbase -> humanized from the local name
    assert tax.label("ifrs-full:CashAndCashEquivalents") == "Cash And Cash Equivalents"


def test_extract_writes_workbook(files, tmp_path):
    out = tmp_path / "out.xlsx"
    result = extract(str(files[0]), schema=str(files[1]), output_xlsx=str(out))
    assert out.exists()
    assert result.fact_count == 4
    assert result.concept_count == 4

    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert {"Overview", "Facts", "Contexts", "Units", "Concepts"} <= set(wb.sheetnames)
    facts = wb["Facts"]
    headers = [c.value for c in facts[1]]
    assert headers[1] == "Concept (QName)" and "Label" in headers
    # 4 data rows
    assert facts.max_row == 5


def test_extract_without_schema_still_works(files):
    result = extract(str(files[0]))  # no schema, no output
    assert result.fact_count == 4
    assert result.taxonomy is None
