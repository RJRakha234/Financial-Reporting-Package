"""Write the extracted XBRL data to a multi-sheet Excel workbook."""

from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .instance import Instance
from .taxonomy import Taxonomy

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, color="1F3864", size=14)
SUB_FONT = Font(bold=True, color="1F3864")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RIGHT = Alignment(horizontal="right")


def _header(ws, headers, row=1):
    """Style an already-appended header row and freeze panes below it."""
    for cell in ws[row]:
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _num(value: str | None):
    """Return an int/float when the raw text is numeric, else the raw string."""
    if value is None:
        return None
    try:
        if "." in value or "e" in value.lower():
            return float(value)
        return int(value)
    except (ValueError, AttributeError):
        return value


def _classify(concept, unit) -> str:
    if concept and concept.type:
        t = concept.type.lower()
        if "monetary" in t:
            return "Monetary"
        if "shares" in t:
            return "Shares"
        if "pershare" in t or "per_share" in t:
            return "Per share"
        if "textblock" in t:
            return "Text block"
        if "string" in t:
            return "String"
        if "date" in t:
            return "Date"
        if "decimal" in t or "pure" in t or "percent" in t:
            return "Numeric"
    if unit:
        m = unit.measure.lower()
        if "/" in m:
            return "Per share"
        if m in ("shares",):
            return "Shares"
        if m in ("pure",):
            return "Numeric"
        return "Monetary"
    return "Non-numeric"


def build_workbook(inst: Instance, tax: Taxonomy | None, source: str) -> Workbook:
    wb = Workbook()
    fact_counts = Counter(f.concept for f in inst.facts)

    _overview_sheet(wb.active, inst, tax, source, fact_counts)
    _facts_sheet(wb.create_sheet("Facts"), inst, tax)
    _contexts_sheet(wb.create_sheet("Contexts"), inst)
    _units_sheet(wb.create_sheet("Units"), inst)
    if tax:
        _concepts_sheet(wb.create_sheet("Concepts"), inst, tax, fact_counts)
        _statements_sheet(wb.create_sheet("Statements"), tax)
        _presentation_sheet(wb.create_sheet("Presentation"), tax)
        _calculation_sheet(wb.create_sheet("Calculation"), tax)
    return wb


def _overview_sheet(ws, inst, tax, source, fact_counts):
    ws.title = "Overview"
    dei = {f.localname: f.value for f in inst.facts if f.prefix == "dei"}
    rows = [
        ("Infosys — extracted XBRL facts", TITLE_FONT),
        ("", None),
        (f"Registrant: {dei.get('EntityRegistrantName', '')}", None),
        (f"Document type: {dei.get('DocumentType', '')}", None),
        (f"Period end: {dei.get('DocumentPeriodEndDate', '')}", None),
        (f"CIK: {dei.get('EntityCentralIndexKey', '')}", None),
        ("", None),
        ("Contents", SUB_FONT),
        (f"  Facts: {len(inst.facts)}", None),
        (f"  Distinct concepts: {len(fact_counts)}", None),
        (f"  Contexts: {len(inst.contexts)}", None),
        (f"  Units: {len(inst.units)}", None),
    ]
    if tax:
        rows += [
            (f"  Presentation roles (statements/disclosures): {len(tax.presentation)}", None),
            (f"  Calculation roles: {len(tax.calculation)}", None),
        ]
    rows += [
        ("", None),
        ("Sheets", SUB_FONT),
        ("  Facts        — one row per tagged value (the core extract)", None),
        ("  Contexts     — every period / dimension context, resolved", None),
        ("  Units        — unit definitions", None),
        ("  Concepts     — every concept used: type, balance, period, label", None),
        ("  Statements   — index of presentation roles", None),
        ("  Presentation — ordered concept outline per statement", None),
        ("  Calculation  — parent = Σ(child × weight) footing relationships", None),
        ("", None),
        (f"Source instance: {source}", None),
    ]
    for i, (text, font) in enumerate(rows, 1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    ws.column_dimensions["A"].width = 90


def _facts_sheet(ws, inst, tax):
    headers = [
        "Fact ID", "Concept (QName)", "Label", "Type", "Value", "Unit",
        "Decimals", "Period", "Start", "End", "Dimensions (Axis = Member)",
        "Negated label", "Context",
    ]
    ws.append(headers)
    for f in inst.facts:
        ctx = inst.context_of(f)
        unit = inst.unit_of(f)
        concept = tax.concepts.get(f.concept) if tax else None
        label = tax.label(f.concept) if tax else f.localname
        negated = "Y" if (concept and concept.negated) else ""
        ws.append([
            f.id,
            f.concept,
            label,
            _classify(concept, unit),
            "(nil)" if f.is_nil else _num(f.value),
            unit.measure if unit else "",
            f.decimals,
            ctx.period_type if ctx else "",
            ctx.start if ctx else "",
            ctx.instant if (ctx and ctx.period_type == "instant") else (ctx.end if ctx else ""),
            ctx.dims_str() if ctx else "",
            negated,
            f.context_ref,
        ])
    _header(ws, headers)
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        row[4].alignment = RIGHT  # Value
        if isinstance(row[4].value, (int, float)):
            row[4].number_format = "#,##0.##"
        for cell in row:
            cell.border = BORDER
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    _widths(ws, [22, 46, 34, 12, 18, 10, 9, 10, 12, 12, 50, 8, 22])


def _contexts_sheet(ws, inst):
    headers = ["Context ID", "Entity (CIK)", "Period", "Start", "End/Instant",
               "Dimensions (Axis = Member)"]
    ws.append(headers)
    for cid, c in inst.contexts.items():
        ws.append([
            cid, c.entity, c.period_type, c.start,
            c.instant if c.period_type == "instant" else c.end, c.dims_str(),
        ])
    _header(ws, headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    _widths(ws, [40, 14, 10, 12, 12, 60])


def _units_sheet(ws, inst):
    headers = ["Unit ID", "Measure", "Raw"]
    ws.append(headers)
    for uid, u in inst.units.items():
        ws.append([uid, u.measure, u.raw])
    _header(ws, headers)
    _widths(ws, [20, 18, 40])


def _concepts_sheet(ws, inst, tax, fact_counts):
    headers = ["Concept (QName)", "Label", "Type", "Period", "Balance",
               "Abstract", "Source", "Negated", "# facts"]
    ws.append(headers)
    used = sorted(fact_counts)
    for q in used:
        c = tax.concepts.get(q)
        ws.append([
            q,
            tax.label(q),
            (c.type if c else ""),
            (c.period_type if c else ""),
            (c.balance if c else ""),
            ("Y" if c and c.abstract else ""),
            ("extension" if c and c.extension else "base"),
            ("Y" if c and c.negated else ""),
            fact_counts[q],
        ])
    _header(ws, headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        row[8].alignment = RIGHT
    _widths(ws, [46, 40, 20, 10, 8, 9, 11, 9, 8])


def _statements_sheet(ws, tax):
    headers = ["No.", "Kind", "Name", "# presentation concepts", "# calc relations"]
    ws.append(headers)
    for uri, role in sorted(tax.roles.items(), key=lambda kv: kv[1].number):
        npres = len(tax.presentation.get(uri, []))
        ncalc = len(tax.calculation.get(uri, []))
        if npres or ncalc:
            ws.append([role.number, role.kind, role.name, npres, ncalc])
    _header(ws, headers)
    _widths(ws, [10, 14, 60, 24, 18])


def _presentation_sheet(ws, tax):
    headers = ["Statement", "Indent", "Concept (QName)", "Label", "Pref. label role"]
    ws.append(headers)
    for uri, nodes in sorted(
        tax.presentation.items(),
        key=lambda kv: tax.roles.get(kv[0]).number if tax.roles.get(kv[0]) else "",
    ):
        role = tax.roles.get(uri)
        name = role.name if role else uri
        for n in nodes:
            pref = n.preferred_label.rsplit("/", 1)[-1] if n.preferred_label else ""
            label = ("    " * n.depth) + tax.label(n.qname)
            ws.append([name, n.depth, n.qname, label, pref])
    _header(ws, headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    _widths(ws, [44, 8, 46, 52, 16])


def _calculation_sheet(ws, tax):
    headers = ["Statement", "Parent (total)", "Child (component)", "Weight"]
    ws.append(headers)
    for uri, rels in sorted(
        tax.calculation.items(),
        key=lambda kv: tax.roles.get(kv[0]).number if tax.roles.get(kv[0]) else "",
    ):
        role = tax.roles.get(uri)
        name = role.name if role else uri
        for r in rels:
            ws.append([name, tax.label(r.parent), tax.label(r.child), r.weight])
    _header(ws, headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    _widths(ws, [44, 46, 46, 8])
