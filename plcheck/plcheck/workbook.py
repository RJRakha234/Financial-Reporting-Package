"""Assemble the ``Check`` workbook with live reconciliation formulas.

The output is a genuine, self-contained Excel file: the three source reports
are embedded as sheets and every difference is a real formula referencing them,
so it recalculates in Excel exactly like a hand-built check file — but produced
in one step from the inputs.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from . import config as C
from . import inputs
from .layout import (COL_ACCOUNT, COL_CATEGORY, COL_CLASS, COL_DESC,
                     CheckLayout, plan)
from .model import ReportTable, SheetData

# Row map of the Check sheet (matches the reference layout).
ROW_TB_LABEL = 2       # "Real Time TB" + per-entity MATCH into the TB
ROW_AGG_LABEL = 3      # "Aggregate Exp" + per-entity MATCH / FX-rate VLOOKUP
ROW_SUMDIFF = 4        # "Sum of Differences" + SUM of each difference column
ROW_BLOCKLABEL = 5
ROW_SUBHEADER = 6
ROW_ENTNAME = 7
ROW_CURR = 8
FIRST_DATA_ROW = 9

RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006")
HDR_FONT = Font(bold=True)
HELP_FONT = Font(italic=True, color="808080")


def _embed(wb: Workbook, sheet: SheetData) -> None:
    ws = wb.create_sheet(sheet.title)
    for coord, value in sheet.cells.items():
        ws[coord] = value
        if coord in sheet.text_coords:
            ws[coord].data_type = "s"      # keep "=..."-looking notes as text


def build_check_workbook(report: ReportTable, tb_path: str, agg_path: str,
                         rates_path: str,
                         cfg: C.CheckConfig | None = None,
                         consol_path: str | None = None) -> Workbook:
    cfg = cfg or C.CheckConfig()
    # LC - Consol is reconciled on its own dedicated sheet (not the Check tab),
    # so the Check tab keeps only the LC-Consol *value* columns (needed for FX).
    layout = plan(report)
    codes = [e.code for e in report.entities]
    currency = {e.code: e.currency for e in report.entities}

    tb = inputs.parse_tb(tb_path, codes)
    agg = inputs.parse_agg(agg_path, codes) if agg_path else None
    rates = inputs.parse_rates(rates_path)
    consol = inputs.parse_consol(consol_path) if consol_path else None
    consol_ref = None
    if consol:
        cc = get_column_letter(consol.concat_col)
        concat_rng = (f"'{C.SHEET_CONSOL}'!${cc}${consol.first_row}"
                      f":${cc}${consol.last_row}")
        val_rngs = [f"'{C.SHEET_CONSOL}'!${get_column_letter(vc)}${consol.first_row}"
                    f":${get_column_letter(vc)}${consol.last_row}"
                    for vc in consol.val_cols]
        func_rng = None
        if consol.func_col:
            fc = get_column_letter(consol.func_col)
            func_rng = (f"'{C.SHEET_CONSOL}'!${fc}${consol.first_row}"
                        f":${fc}${consol.last_row}")
        consol_ref = (concat_rng, val_rngs, func_rng)
    gc_currency = cfg.gc_currency or inputs.detect_gc_currency(report, rates.rates)

    tb_lastcol = get_column_letter(tb.last_col)
    agg_mfirst = get_column_letter(agg.match_first_col) if agg else None
    agg_mlast = get_column_letter(agg.match_last_col) if agg else None
    rate_from = get_column_letter(rates.from_col)
    rate_to = get_column_letter(rates.rate_col)

    # Span the FX / consolidation source ranges over whichever of the expected
    # blocks are actually present, so a renamed or absent block can't crash the
    # build — the missing block simply contributes nothing to the SUMIF.
    def _first_present(blocks):
        cols = [layout.block_first_value[b] for b in blocks
                if b in layout.block_first_value]
        return min(cols, key=column_index_from_string) if cols else None

    def _last_present(blocks):
        cols = [layout.block_last_value[b] for b in blocks
                if b in layout.block_last_value]
        return max(cols, key=column_index_from_string) if cols else None

    fx_last = _last_present(C.FX_SOURCE_BLOCKS)
    consol_first = _first_present(C.CONSOL_SOURCE_BLOCKS)
    consol_last = _last_present(C.CONSOL_SOURCE_BLOCKS)

    # MA-rate lookup spans a safe fixed window so extra currencies are covered.
    rate_last_row = max(rates.last_row,
                        rates.first_row + cfg.rate_lookup_rows - 1)

    wb = Workbook()
    ws = wb.active
    ws.title = C.SHEET_CHECK

    # ---- header band ------------------------------------------------------
    ws[f"{COL_ACCOUNT}{ROW_CURR}"] = "GLACCOUNT"
    ws[f"{COL_DESC}{ROW_CURR}"] = "GL Description | Currency"
    ws[f"{COL_DESC}{ROW_TB_LABEL}"] = C.SHEET_TB
    ws[f"{COL_DESC}{ROW_AGG_LABEL}"] = C.SHEET_AGG
    ws[f"{COL_DESC}{ROW_SUMDIFF}"] = "Sum of Differences"
    for cell in (f"{COL_DESC}{ROW_TB_LABEL}", f"{COL_DESC}{ROW_AGG_LABEL}",
                 f"{COL_DESC}{ROW_SUMDIFF}"):
        ws[cell].font = HELP_FONT

    ent_by_valuecol: dict[str, str] = {}      # LC-Balance value col -> code
    for info in layout.cols:
        ws[f"{info.value_col}{ROW_BLOCKLABEL}"] = info.block
        ws[f"{info.value_col}{ROW_SUBHEADER}"] = info.sub
        ws[f"{info.value_col}{ROW_BLOCKLABEL}"].font = HDR_FONT
        if info.is_entity:
            e = next(en for en in report.entities if en.code == info.sub)
            ws[f"{info.value_col}{ROW_ENTNAME}"] = e.name
            ws[f"{info.value_col}{ROW_CURR}"] = e.currency
            if info.block == C.CHECK_LC_BALANCE:
                ent_by_valuecol[info.value_col] = e.code

    # ---- per-entity helper cells -----------------------------------------
    for info in layout.cols:
        if not (info.is_entity and info.diff_col):
            continue
        d, v = info.diff_col, info.value_col
        if info.block == C.CHECK_LC_BALANCE:
            ws[f"{d}{ROW_TB_LABEL}"] = (
                f"=MATCH({v}${ROW_SUBHEADER},'{C.SHEET_TB}'!"
                f"$A${tb.header_row}:${tb_lastcol}${tb.header_row},0)")
            if agg:
                ws[f"{d}{ROW_AGG_LABEL}"] = (
                    f"=MATCH({v}${ROW_SUBHEADER},'{C.SHEET_AGG}'!"
                    f"${agg_mfirst}${agg.subheader_row}:${agg_mlast}${agg.subheader_row},0)")
        elif info.block == C.CHECK_GC_BALANCE:
            rng = (f"${rate_from}${rates.first_row}:"
                   f"${rate_to}${rate_last_row}")
            local = (f"VLOOKUP({v}${ROW_CURR},'{C.SHEET_RATES}'!{rng},"
                     f"{rates.col_index},FALSE)")
            # The rates table quotes every currency TO the base (e.g. INR). If
            # the group currency is itself the base, it has no row of its own,
            # so the rate already converts to it (divisor = 1). Only when the
            # group currency is a quoted currency (e.g. USD) do we cross-divide.
            if gc_currency in rates.rates:
                ws[f"{d}{ROW_AGG_LABEL}"] = (
                    f"={local}/VLOOKUP(\"{gc_currency}\",'{C.SHEET_RATES}'!{rng},"
                    f"{rates.col_index},FALSE)")
            else:
                ws[f"{d}{ROW_AGG_LABEL}"] = f"={local}"
        for hr in (ROW_TB_LABEL, ROW_AGG_LABEL):
            if ws[f"{d}{hr}"].value is not None:
                ws[f"{d}{hr}"].font = HELP_FONT

    # ---- data rows --------------------------------------------------------
    diff_cols: set[str] = set()
    last_data_row = FIRST_DATA_ROW - 1
    for i, row in enumerate(report.rows):
        r = FIRST_DATA_ROW + i
        last_data_row = r
        rule = cfg.effective_rule(row.category)
        ws[f"{COL_CATEGORY}{r}"] = row.category or None
        if not row.is_subtotal:
            if rule and rule.classification:
                ws[f"{COL_CLASS}{r}"] = rule.classification
            ws[f"{COL_ACCOUNT}{r}"] = row.account
            ws[f"{COL_DESC}{r}"] = row.description

        # value cells (verbatim from the report)
        for info in layout.cols:
            val = row.values.get((info.block, info.sub))
            if val is not None:
                ws[f"{info.value_col}{r}"] = val

        # difference formulas
        for info in layout.cols:
            if not (info.is_entity and info.diff_col):
                continue
            d, v = info.diff_col, info.value_col
            diff_cols.add(d)
            f = _diff_formula(info, r, d, v, rule, agg, tb, layout,
                              fx_last, consol_first, consol_last,
                              row.is_subtotal)
            if f:
                ws[f"{d}{r}"] = f

    # ---- sum-of-differences row ------------------------------------------
    for d in sorted(diff_cols):
        ws[f"{d}{ROW_SUMDIFF}"] = f"=SUM({d}{FIRST_DATA_ROW}:{d}{last_data_row})"

    # ---- net-profit reconciliation block ---------------------------------
    _net_profit_block(ws, report, layout, tb, last_data_row, cfg)

    # ---- highlighting + cosmetics ----------------------------------------
    _highlight(ws, diff_cols, FIRST_DATA_ROW, last_data_row, cfg.tolerance)
    _cosmetics(ws, layout)

    # ---- derived Minority Interest sheet ---------------------------------
    _minority_sheet(wb, report, layout, cfg)

    # ---- dedicated LC - Consol check sheet (tie-out to the tracker) ------
    if consol_ref is not None:
        _lc_consol_sheet(wb, report, layout, consol, consol_ref, cfg)

    # ---- entity (company-code) coverage across the inputs ----------------
    _entity_coverage_sheet(wb, report, tb_path, agg_path)

    # ---- selected GLs verified against the TB (independent tab) -----------
    _tb_override_sheet(wb, report, layout, tb, cfg)

    # ---- embed the source sheets -----------------------------------------
    _embed(wb, inputs.read_sheet(tb_path, C.SHEET_TB))
    if agg_path:
        _embed(wb, inputs.read_sheet(agg_path, C.SHEET_AGG))
    _embed(wb, inputs.read_sheet(rates_path, C.SHEET_RATES))
    if consol_path:
        # values-only: a manual tracker often carries formulas that point at
        # its *other* tabs, which would arrive broken (and make Excel offer to
        # "repair" the file); the tie-out SUMIFs only need the values anyway.
        _embed(wb, inputs.read_sheet(consol_path, C.SHEET_CONSOL,
                                     keep_formulas=False,
                                     sheet_name=consol.sheet_name))
        # mirror the entry-level functional-code fill onto the embedded copy, so
        # the live SUMIFS (which matches the code on the P&L leg's row) works
        # even when the code was typed on the contra leg.
        if consol.func_col:
            inputs.propagate_func_column(
                wb[C.SHEET_CONSOL], consol.concat_col, consol.func_col,
                consol.val_cols[0],
                consol.val_cols[1] if len(consol.val_cols) > 1 else 0,
                consol.first_row, consol.last_row)
    return wb


MINORITY_SHEET = "Minority Interest"


def _minority_sheet(ws_wb, report: ReportTable, layout: CheckLayout,
                    cfg: C.CheckConfig) -> None:
    """Add the Minority Interest sheet, derived from the Check sheet.

    Per company code (every entity in the LC section):
      Net Profit as per GC Bal  = GC-Balance of the Net Profit line
      Div received - <acct>     = GC-Balance of the dividend GL (0 if absent)
      Profit before Div         = the two added
      Minority - Total          = GC-Total of the Minority Interest line
      Current Period %          = Minority-Total / Profit-before-Div
    """
    # needs the GC-Balance and GC-Total blocks, plus Net Profit / Minority lines
    if (C.CHECK_GC_BALANCE not in layout.block_first_value
            or C.CHECK_GC_TOTAL not in layout.block_first_value):
        return

    def _norm(s):
        return str(s).strip().lower()

    def section_row(pred):
        idxs = [i for i, r in enumerate(report.rows) if pred(r.category)]
        if not idxs:
            return None
        subs = [i for i in idxs if report.rows[i].is_subtotal]
        return FIRST_DATA_ROW + (subs[0] if subs else idxs[0])

    np_row = section_row(lambda c: _norm(c) == "net profit")
    # fuzzy so a misspelled "Minority Interest" section still feeds the sheet
    mi_row = section_row(cfg.is_minority)
    if np_row is None or mi_row is None:
        return
    div_idx = next((i for i, r in enumerate(report.rows)
                    if inputs.account_key(r.account) == cfg.dividend_account),
                   None)
    div_row = FIRST_DATA_ROW + div_idx if div_idx is not None else None

    ws = ws_wb.create_sheet(MINORITY_SHEET)
    headers = ["Code", "Net Profit as per GC Bal",
               f"Div received - {cfg.dividend_account}", "Profit before Div",
               "Minority - Total", "Current Period %"]
    for j, h in enumerate(headers, start=1):
        ws.cell(1, j, h).font = HDR_FONT

    chk = C.SHEET_CHECK
    for k, e in enumerate(report.entities):
        row = 2 + k
        gcbal = layout.value_col(C.CHECK_GC_BALANCE, e.code)
        gctot = layout.value_col(C.CHECK_GC_TOTAL, e.code)
        ws.cell(row, 1, e.code)
        ws.cell(row, 2).value = f"='{chk}'!{gcbal}{np_row}"
        ws.cell(row, 3).value = (f"='{chk}'!{gcbal}{div_row}"
                                 if div_row else 0)
        ws.cell(row, 4).value = f"=B{row}+C{row}"
        ws.cell(row, 5).value = f"='{chk}'!{gctot}{mi_row}"
        ws.cell(row, 6).value = f"=IFERROR(E{row}/D{row},0)"
        ws.cell(row, 6).number_format = "0.00%"

    for col, w in {"A": 12, "B": 24, "C": 22, "D": 18, "E": 18,
                   "F": 16}.items():
        ws.column_dimensions[col].width = w


def _consol_sum(consol_ref, key, fcode) -> str:
    """The tracker net posting (Debit - Credit) for one company+account, as an
    Excel expression. Matches account-only (nature-wise) or account + Functional
    Group (function-wise split, via SUMIFS)."""
    concat_rng, val_rngs, func_rng = consol_ref
    if fcode and func_rng:
        def _s(vr):
            return f'SUMIFS({vr},{concat_rng},{key},{func_rng},"{fcode}")'
    else:
        def _s(vr):
            return f"SUMIF({concat_rng},{key},{vr})"
    terms = _s(val_rngs[0])
    if len(val_rngs) > 1:
        terms += f"-{_s(val_rngs[1])}"
    return terms


def _diff_formula(info, r, d, v, rule, agg, tb, layout: CheckLayout,
                  fx_last, consol_first, consol_last,
                  is_subtotal=False) -> str | None:
    """The difference formula for one cell, by block type."""
    if info.block == C.CHECK_LC_BALANCE:
        # subtotal rows have no GL account to look up — no LC tie-out
        if is_subtotal or rule is None or not rule.source:
            return None
        match = f"{C.SHEET_CHECK}!{d}${ROW_TB_LABEL}"
        if rule.source == "tb":
            table = (f"'{C.SHEET_TB}'!$A${tb.header_row}:"
                     f"${get_column_letter(tb.last_col)}${tb.last_data_row}")
            return f"=IFERROR(VLOOKUP($C{r},{table},{match},FALSE),0)-{v}{r}"
        blk = agg.blocks.get(rule.source) if agg else None
        if blk is None:
            return None
        match = f"{C.SHEET_CHECK}!{d}${ROW_AGG_LABEL}"
        table = (f"'{C.SHEET_AGG}'!${get_column_letter(blk.acct_col)}"
                 f"${blk.subheader_row}:${get_column_letter(blk.last_entity_col)}"
                 f"${blk.last_data_row}")
        return f"=IFERROR(VLOOKUP($C{r},{table},{match},FALSE),0)-{v}{r}"

    if info.block == C.CHECK_GC_BALANCE:
        if fx_last is None:
            return None
        rate = f"{C.SHEET_CHECK}!{d}${ROW_AGG_LABEL}"
        rng_hdr = f"${COL_CLASS}${ROW_SUBHEADER}:${fx_last}${ROW_SUBHEADER}"
        rng_row = f"${COL_CLASS}{r}:${fx_last}{r}"
        return (f"=SUMIF({rng_hdr},{v}${ROW_SUBHEADER},{rng_row})*{rate}-{v}{r}")

    if info.block == C.CHECK_GC_TOTAL:
        if consol_first is None or consol_last is None:
            return None
        rng_hdr = f"${consol_first}${ROW_SUBHEADER}:${consol_last}${ROW_SUBHEADER}"
        rng_row = f"${consol_first}{r}:${consol_last}{r}"
        return f"=SUMIF({rng_hdr},{v}${ROW_SUBHEADER},{rng_row})-{v}{r}"
    return None


COVERAGE_SHEET = "Entity Coverage"


def _entity_coverage_sheet(wbk, report: ReportTable, tb_path, agg_path) -> None:
    """List the company codes in each input and flag any that are present in one
    source but missing from another (e.g. a TB-only code like a typo)."""
    pl_codes = [e.code for e in report.entities
                if inputs.looks_like_company_code(e.code)]
    sources = [("PL Report", pl_codes),
               ("Real Time TB", inputs.company_codes(tb_path))]
    if agg_path:
        sources.append(("Aggregate Exp", inputs.company_codes(agg_path)))
    sets = {name: {c.upper() for c in codes} for name, codes in sources}

    ws = wbk.create_sheet(COVERAGE_SHEET)
    ws["A1"] = "Company codes by source"
    ws["A1"].font = HDR_FONT
    # raw list per source (one column each)
    for j, (name, codes) in enumerate(sources, start=1):
        ws.cell(2, j, name).font = HDR_FONT
        for k, code in enumerate(codes, start=3):
            ws.cell(k, j, code)

    # union reconciliation matrix
    base = len(sources) + 2
    ws.cell(2, base, "Company Code").font = HDR_FONT
    for j, (name, _c) in enumerate(sources):
        ws.cell(2, base + 1 + j, f"In {name}").font = HDR_FONT
    note_col = base + 1 + len(sources)
    ws.cell(2, note_col, "Note (where missing)").font = HDR_FONT

    union, seen = [], set()
    for _name, codes in sources:
        for c in codes:
            if c.upper() not in seen:
                seen.add(c.upper())
                union.append(c)

    row = 3
    for code in union:
        present = [name for name, _c in sources if code.upper() in sets[name]]
        missing = [name for name, _c in sources if code.upper() not in sets[name]]
        ws.cell(row, base, code)
        for j, (name, _c) in enumerate(sources):
            ws.cell(row, base + 1 + j, "Yes" if name in present else "-")
        note = "in all sources" if not missing else "missing from: " + ", ".join(missing)
        nc = ws.cell(row, note_col, note)
        if missing:
            for col in range(base, note_col + 1):
                ws.cell(row, col).fill = RED_FILL
            nc.font = RED_FONT
        row += 1

    # overall banner
    bad = sum(1 for code in union
              if any(code.upper() not in sets[name] for name, _c in sources))
    ws["A1"] = ("Entity coverage: ALL company codes present in every source"
                if bad == 0 else
                f"Entity coverage: {bad} code(s) missing from some source - see red")
    ws["A1"].font = HDR_FONT if bad == 0 else Font(bold=True, color="9C0006")
    for col, w in {"A": 16, "B": 16, "C": 16}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[get_column_letter(note_col)].width = 30


LC_CONSOL_SHEET = "LC-Consol Check"


def _lc_consol_sheet(wbk, report: ReportTable, layout: CheckLayout,
                     consol, consol_ref, cfg: C.CheckConfig) -> None:
    """Dedicated LC - Consol tie-out sheet. One row per company + P&L account
    *that has a balance* (a report LC-Consol figure or a tracker entry), so only
    the lines with consolidation activity are listed - not every P&L account.
    Each row shows the report value vs the tracker net (Debit - Credit;
    function-wise rows match their COS/S&M/G&A slice), with a tie banner + red."""
    ws = wbk.create_sheet(LC_CONSOL_SHEET)
    tol = cfg.tolerance

    heads = ["Company", "Section", "GLACCOUNT", "GL Description",
             "LC-Consol value", "Diff"]
    for j, h in enumerate(heads, start=1):
        ws.cell(2, j, h).font = HDR_FONT

    first = 3
    k = first
    for i, row in enumerate(report.rows):
        if row.is_subtotal or row.account is None or not cfg.is_pl_account(row.account):
            continue
        check_row = FIRST_DATA_ROW + i
        fcode = cfg.functional_code(row.category)
        for e in report.entities:
            rep_val = row.values.get((C.CHECK_LC_CONSOL, e.code), 0.0) or 0.0
            key_t = inputs.consol_key(e.code, row.account)
            if fcode and consol.func_col:
                trk = consol.totals_fn.get((key_t, fcode.upper()), 0.0)
            else:
                trk = consol.totals.get(key_t, 0.0)
            # show only lines with a balance on either side
            if abs(rep_val) <= tol and abs(trk) <= tol:
                continue
            lc_valcol = layout.value_col(C.CHECK_LC_CONSOL, e.code)
            ws.cell(k, 1, e.code)
            ws.cell(k, 2, row.category or None)
            ws.cell(k, 3, row.account)
            ws.cell(k, 4, row.description)
            ws.cell(k, 5).value = f"='{C.SHEET_CHECK}'!{lc_valcol}{check_row}"
            ws.cell(k, 5).number_format = "#,##0.00"
            key = f'"{e.code}"&$C{k}'
            ws.cell(k, 6).value = f"={_consol_sum(consol_ref, key, fcode)}-E{k}"
            ws.cell(k, 6).number_format = "#,##0.00"
            k += 1
    last = k - 1

    if last >= first:
        ws.conditional_formatting.add(
            f"F{first}:F{last}",
            FormulaRule(formula=[f"ABS(F{first})>{tol}"],
                        fill=RED_FILL, font=RED_FONT))
        ws["B1"] = f"=SUMPRODUCT(--(ABS(F{first}:F{last})>{tol}))"
        ws["A1"] = ('=IF(B1=0,"LC - Consol: ALL entries tie",'
                    '"LC - Consol: "&B1&" difference(s) NOT tied - see red")')
        ws["A1"].font = HDR_FONT
        ws.conditional_formatting.add(
            "A1:B1", FormulaRule(formula=["$B$1>0"], fill=RED_FILL, font=RED_FONT))
    else:
        ws["A1"] = "LC - Consol: no consolidation entries with a balance"
        ws["A1"].font = HDR_FONT

    # --- completeness: P&L consol entries NOT reflected in the report ------
    # (the reverse direction: a tracker entry whose GL isn't a line in the
    # report is never summed by the check above, i.e. left unreconciled.)
    report_codes = {e.code.upper() for e in report.entities}
    report_pl = {inputs.account_key(r.account) for r in report.rows
                 if not r.is_subtotal and r.account is not None
                 and cfg.is_pl_account(r.account)}
    orphans = []
    for key, (comp, acct, desc) in consol.key_info.items():
        if not acct or not cfg.is_pl_account(acct):
            continue                                # only P&L-series legs
        net = consol.totals.get(key, 0.0)
        if abs(net) <= tol:
            continue                                # no current-month value
        covered = (inputs.account_key(acct) in report_pl
                   and (not comp or comp.upper() in report_codes))
        if not covered:
            orphans.append((comp, acct, desc, net))

    base = (last + 3) if last >= first else 4
    ws.cell(base, 1, "Unreconciled consol entries (current month, P&L, "
                     "not found in report)").font = HDR_FONT
    if orphans:
        ws.cell(base, 5, f"{len(orphans)} NOT reconciled").font = RED_FONT
    hr = base + 1
    for j, h in enumerate(["Company", "GLACCOUNT", "GL Description",
                           "Net (Dr-Cr)", "Status"], start=1):
        ws.cell(hr, j, h).font = HDR_FONT
    if orphans:
        rr = hr + 1
        for comp, acct, desc, net in orphans:
            ws.cell(rr, 1, comp)
            ws.cell(rr, 2, acct)
            ws.cell(rr, 3, desc)
            ws.cell(rr, 4, round(net, 2)).number_format = "#,##0.00"
            ws.cell(rr, 5, "NOT reconciled")
            for col in range(1, 6):
                ws.cell(rr, col).fill = RED_FILL
                ws.cell(rr, col).font = RED_FONT
            rr += 1
    else:
        ws.cell(hr + 1, 1,
                "All consol entries reconciled.").font = Font(bold=True,
                                                              color="006100")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 14
    ws.freeze_panes = "A3"


TB_OVERRIDE_SHEET = "Selected GL TB Check"


def _tb_override_sheet(ws_wb, report: ReportTable, layout: CheckLayout, tb,
                       cfg: C.CheckConfig) -> None:
    """Independent tab: the configured GL accounts (tb_override_accounts) have
    their report LC-Balance figure verified against the Real Time TB, instead of
    the Aggregate Expenses report their section would normally use. One row per
    company + account; Diff = TB − report, with a banner and red highlighting."""
    wanted = {inputs.account_key(a) for a in cfg.tb_override_accounts}
    if not wanted:
        return
    ws = ws_wb.create_sheet(TB_OVERRIDE_SHEET)
    tol = cfg.tolerance
    heads = ["Company", "GLACCOUNT", "GL Description", "Report LC value",
             "Real Time TB", "Diff"]
    for j, h in enumerate(heads, start=1):
        ws.cell(2, j, h).font = HDR_FONT

    tb_last = get_column_letter(tb.last_col)
    tb_rng = f"'{C.SHEET_TB}'!$A${tb.header_row}:${tb_last}${tb.last_data_row}"
    tb_hdr = f"'{C.SHEET_TB}'!$A${tb.header_row}:${tb_last}${tb.header_row}"

    first = 3
    k = first
    for i, row in enumerate(report.rows):
        if row.is_subtotal or row.account is None:
            continue
        if inputs.account_key(row.account) not in wanted:
            continue
        check_row = FIRST_DATA_ROW + i
        for e in report.entities:
            lc_valcol = layout.value_col(C.CHECK_LC_BALANCE, e.code)
            ws.cell(k, 1, e.code)
            ws.cell(k, 2, row.account)
            ws.cell(k, 3, row.description)
            ws.cell(k, 4).value = f"='{C.SHEET_CHECK}'!{lc_valcol}{check_row}"
            ws.cell(k, 5).value = (f'=IFERROR(VLOOKUP($B{k},{tb_rng},'
                                   f'MATCH("{e.code}",{tb_hdr},0),FALSE),0)')
            ws.cell(k, 6).value = f"=E{k}-D{k}"
            for col in (4, 5, 6):
                ws.cell(k, col).number_format = "#,##0.00"
            k += 1
    last = k - 1

    if last >= first:
        ws.conditional_formatting.add(
            f"F{first}:F{last}",
            FormulaRule(formula=[f"ABS(F{first})>{tol}"],
                        fill=RED_FILL, font=RED_FONT))
        ws["B1"] = f"=SUMPRODUCT(--(ABS(F{first}:F{last})>{tol}))"
        ws["A1"] = ('=IF(B1=0,"Selected GLs: ALL tie to TB",'
                    '"Selected GLs: "&B1&" difference(s) vs TB - see red")')
        ws["A1"].font = HDR_FONT
        ws.conditional_formatting.add(
            "A1:B1", FormulaRule(formula=["$B$1>0"], fill=RED_FILL, font=RED_FONT))
    else:
        ws["A1"] = ("Selected GL TB Check: none of the configured accounts "
                    f"({', '.join(str(a) for a in cfg.tb_override_accounts)}) "
                    "are in this report")
        ws["A1"].font = HDR_FONT
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 32
    for col in ("D", "E", "F"):
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A3"


def _net_profit_block(ws, report: ReportTable, layout: CheckLayout, tb,
                      last_data_row, cfg: C.CheckConfig) -> None:
    start = last_data_row + 2
    r_inc, r_exp, r_np = start, start + 1, start + 2
    r_calc = start + 4
    r_tbnp = start + 6
    r_chk = start + 7
    ws[f"{COL_DESC}{r_inc}"] = "Income"
    ws[f"{COL_DESC}{r_exp}"] = "Expense"
    ws[f"{COL_DESC}{r_np}"] = "Net Profit"
    ws[f"{COL_DESC}{r_calc}"] = "Calc Check"
    ws[f"{COL_DESC}{r_tbnp}"] = "Net Profit as per Real Time TB"
    ws[f"{COL_DESC}{r_chk}"] = "Check"
    # TB net profit is summed live from the P&L-series GLs (leading digit 1/2/3)
    # using a text-safe SUMPRODUCT, so it works whether the TB stores account
    # numbers as numbers or as text, and tracks GLs being added / removed.
    acct = get_column_letter(tb.acct_col)
    first, last = tb.first_data_row, tb.last_data_row
    acct_rng = f"'{C.SHEET_TB}'!${acct}${first}:${acct}${last}"
    series_test = "+".join(
        f'(LEFT(TRIM({acct_rng}),1)="{d}")' for d in cfg.pl_series)
    for e in report.entities:
        v = layout.value_col(C.CHECK_LC_BALANCE, e.code)   # value column (E/G/I)
        d = layout.diff_col(C.CHECK_LC_BALANCE, e.code)    # output column (F/H/J)
        ws[f"{d}{r_inc}"] = f"=SUMIF($A:$A,$D{r_inc},{v}:{v})"
        ws[f"{d}{r_exp}"] = f"=SUMIF($A:$A,$D{r_exp},{v}:{v})"
        ws[f"{d}{r_np}"] = f"=SUMIF($A:$A,$D{r_np},{v}:{v})"
        ws[f"{d}{r_calc}"] = f"={d}{r_inc}+{d}{r_exp}+{d}{r_np}"
        # Only sum the TB when this company actually exists there; otherwise the
        # net profit per TB is undefined (0), not the first company's figure.
        if e.code in tb.entity_cols:
            ent = get_column_letter(tb.entity_cols[e.code])
            ent_rng = f"'{C.SHEET_TB}'!${ent}${first}:${ent}${last}"
            ws[f"{d}{r_tbnp}"] = f"=SUMPRODUCT(({series_test})*({ent_rng}))"
        else:
            ws[f"{d}{r_tbnp}"] = 0
        ws[f"{d}{r_chk}"] = f"={d}{r_np}+{d}{r_tbnp}"
    for rr in (r_inc, r_exp, r_np, r_calc, r_tbnp, r_chk):
        ws[f"{COL_DESC}{rr}"].font = HDR_FONT


def _highlight(ws, diff_cols, first_row, last_row, tol) -> None:
    # data rows, the Sum-of-Differences row, and the bottom block share the
    # difference columns; flag any cell whose absolute value exceeds tolerance.
    for d in diff_cols:
        top = f"{d}{ROW_SUMDIFF}"
        rng = f"{d}{ROW_SUMDIFF}:{d}{last_row + 12}"
        # the formula is relative to the top-left of the range, so it adjusts
        # row-by-row down the column.
        rule = FormulaRule(formula=[f"AND({top}<>\"\",ABS({top})>{tol})"],
                           fill=RED_FILL, font=RED_FONT)
        ws.conditional_formatting.add(rng, rule)


def _cosmetics(ws, layout: CheckLayout) -> None:
    ws.freeze_panes = f"{get_column_letter(5)}{FIRST_DATA_ROW}"
    widths = {COL_CLASS: 16, COL_CATEGORY: 22, COL_ACCOUNT: 11, COL_DESC: 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for info in layout.cols:
        ws.column_dimensions[info.value_col].width = 15
        if info.diff_col:
            ws.column_dimensions[info.diff_col].width = 11
            ws[f"{info.diff_col}{ROW_BLOCKLABEL}"] = "Diff"
            ws[f"{info.diff_col}{ROW_BLOCKLABEL}"].font = HELP_FONT
