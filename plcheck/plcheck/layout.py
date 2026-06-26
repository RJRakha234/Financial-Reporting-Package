"""Plan the column layout of the generated ``Check`` sheet.

The check sheet mirrors the report's blocks but inserts a *difference* column
after every entity column of the three reconciled blocks (LC - Balance,
GC - Balance, GC - Total).  This module works out, once, which letter each
value and difference column lands on so the builder and the evaluator agree.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from . import config as C
from .model import ReportTable

# Label columns at the far left of the check sheet.
COL_CLASS = "A"      # Income / Expense / Net Profit
COL_CATEGORY = "B"   # Revenue, Cost of Production, ...
COL_ACCOUNT = "C"    # GLACCOUNT
COL_DESC = "D"       # GL Description | Currency
FIRST_VALUE_COL = 5  # column E

CHECKED_BLOCKS = {C.CHECK_LC_BALANCE, C.CHECK_GC_BALANCE, C.CHECK_GC_TOTAL}


@dataclass
class ColInfo:
    block: str
    sub: str            # entity code or "Overall Result"
    value_col: str
    diff_col: str | None = None
    is_entity: bool = False


@dataclass
class CheckLayout:
    cols: list[ColInfo] = field(default_factory=list)
    # quick lookups
    value_of: dict[tuple[str, str], str] = field(default_factory=dict)
    diff_of: dict[tuple[str, str], str] = field(default_factory=dict)
    block_first_value: dict[str, str] = field(default_factory=dict)
    block_last_value: dict[str, str] = field(default_factory=dict)

    def value_col(self, block: str, sub: str) -> str:
        return self.value_of[(block, sub)]

    def diff_col(self, block: str, sub: str) -> str | None:
        return self.diff_of.get((block, sub))


def plan(report: ReportTable, extra_checked=()) -> CheckLayout:
    checked_blocks = CHECKED_BLOCKS | set(extra_checked)
    layout = CheckLayout()
    idx = FIRST_VALUE_COL
    for block in report.blocks:
        checked = block.label in checked_blocks
        first_val = None
        last_val = None
        # preserve the report's column order within the block
        for sub in block.columns:
            is_entity = sub != C.OVERALL_LABEL
            value_col = get_column_letter(idx)
            idx += 1
            diff_col = None
            if checked and is_entity:
                diff_col = get_column_letter(idx)
                idx += 1
            info = ColInfo(block=block.label, sub=sub, value_col=value_col,
                           diff_col=diff_col, is_entity=is_entity)
            layout.cols.append(info)
            layout.value_of[(block.label, sub)] = value_col
            if diff_col:
                layout.diff_of[(block.label, sub)] = diff_col
            first_val = first_val or value_col
            last_val = value_col
        layout.block_first_value[block.label] = first_val
        layout.block_last_value[block.label] = last_val
    return layout
