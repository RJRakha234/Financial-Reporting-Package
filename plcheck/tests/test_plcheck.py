"""Tests for plcheck, driven by the bundled sample inputs.

The sample report contains two genuine discrepancies (the General
Administration / Depreciation line was entered as rounded integers, 458 and
157384, instead of 458.51 and 157384.65), which the check must surface — and
which the reference check file also flags.
"""

import os

import openpyxl
import pytest

from plcheck import analyze
from plcheck import inputs
from plcheck.evaluate import evaluate
from plcheck.report import to_dict

HERE = os.path.dirname(__file__)
SAMPLE = os.path.join(HERE, os.pardir, "sample")
REPORT = os.path.join(SAMPLE, "IFRS_INR_PL_Report.xlsx")
TB = os.path.join(SAMPLE, "Real_Time_TB.xlsx")
AGG = os.path.join(SAMPLE, "Aggregate_Expenses.xlsx")
RATES = os.path.join(SAMPLE, "MA_Rates.xlsx")


@pytest.fixture(scope="module")
def report():
    return inputs.read_report(REPORT)


def test_report_parsed(report):
    assert [e.code for e in report.entities] == ["BALSCH", "BALSDE", "BALSDK"]
    assert [e.currency for e in report.entities] == ["CHF", "EUR", "DKK"]
    labels = [b.label for b in report.blocks]
    assert labels[0] == "LC - Balance"
    assert "GC - Total" in labels


def test_evaluation_flags_depreciation_rounding(report):
    ev = evaluate(report, TB, AGG, RATES)
    assert not ev.ok
    # the only LC tie-out failures are the depreciation roundings
    lc = {(d.entity, round(d.delta, 2)) for d in ev.flagged() if d.kind == "lc"}
    assert ("BALSDE", 0.51) in lc
    assert ("BALSDK", 0.65) in lc
    # no LC tie-out failure for any other account
    bad_accounts = {d.account for d in ev.flagged()
                    if d.kind == "lc" and abs(d.delta) > 0.5}
    assert bad_accounts == {290100}


def test_fx_and_consolidation_otherwise_tie(report):
    ev = evaluate(report, TB, AGG, RATES)
    # every consolidation (GC-Total) difference reconciles
    consol = [d for d in ev.flagged() if d.kind == "consol"]
    assert consol == []
    # the only FX failures are the ones propagated from the rounding
    fx = {round(d.delta, 2) for d in ev.flagged() if d.kind == "fx"}
    assert fx <= {-54.99, 54.99, -9.24, 9.23, 9.24}


def test_net_profit_reconciliation(report):
    ev = evaluate(report, TB, AGG, RATES)
    for n in ev.net_profit:
        # internal P&L integrity always holds
        assert abs(n.calc_check) < 0.01
    tie = {n.entity: round(n.tie_check, 2) for n in ev.net_profit}
    assert tie["BALSCH"] == 0.0
    assert tie["BALSDE"] == 0.51
    assert tie["BALSDK"] == 0.65


def test_all_categories_mapped(report):
    ev = evaluate(report, TB, AGG, RATES)
    assert ev.unmapped_categories == []


def test_build_writes_valid_workbook(tmp_path):
    out = tmp_path / "Check.xlsx"
    result = analyze(REPORT, TB, AGG, RATES, output_path=str(out))
    assert result.output_path == str(out)
    assert out.is_file()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Check", "Minority Interest", "Entity Coverage",
                             "Real Time TB", "Aggregate Exp", "MA rates"]
    ws = wb["Check"]
    # difference formulas are present on detail rows, absent on subtotal rows
    assert str(ws["F9"].value).startswith("=IFERROR(VLOOKUP")
    assert ws["F11"].value is None            # Revenue subtotal: no LC tie-out
    assert str(ws["Q9"].value).startswith("=SUMIF")   # FX conversion
    assert str(ws["AJ9"].value).startswith("=SUMIF")  # consolidation


def test_json_summary_shape(report):
    ev = evaluate(report, TB, AGG, RATES)
    d = to_dict(ev)
    assert set(d) >= {"ok", "tolerance", "differences", "net_profit"}
    assert d["ok"] is False
    assert all("delta" in row for row in d["differences"])


# --------------------------------------------------------------------------
# Requirement #1: a different number of GL accounts must never be missed.
# --------------------------------------------------------------------------

def test_added_gl_account_gets_all_checks(tmp_path):
    """A GL added to the report must receive LC / FX / consolidation checks."""
    from plcheck.model import ReportRow
    from plcheck.workbook import build_check_workbook, FIRST_DATA_ROW

    report = inputs.read_report(REPORT)
    # insert a brand-new Cost of Production GL right after the existing ones
    new_row = ReportRow(
        category="Cost of Production", account=110999,
        description="Newly added salary GL",
        values={(b.label, sub): 123.0
                for b in report.blocks for sub in b.columns},
    )
    pos = max(i for i, r in enumerate(report.rows)
              if r.category == "Cost of Production") + 1
    report.rows.insert(pos, new_row)

    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)["Check"]
    r = FIRST_DATA_ROW + pos
    assert ws[f"C{r}"].value == 110999                         # GL carried over
    assert str(ws[f"F{r}"].value).startswith("=IFERROR(VLOOKUP")  # LC tie-out
    assert str(ws[f"Q{r}"].value).startswith("=SUMIF")          # FX
    assert str(ws[f"AJ{r}"].value).startswith("=SUMIF")         # consolidation


def test_lc_checks_cover_every_sourced_detail_gl():
    """Number of LC tie-out checks == sourced detail GLs x entities."""
    from plcheck.config import CheckConfig
    report = inputs.read_report(REPORT)
    cfg = CheckConfig()
    sourced = [r for r in report.rows
               if not r.is_subtotal and r.category
               and (cfg.rule_for(r.category) and cfg.rule_for(r.category).source)]
    ev = evaluate(report, TB, AGG, RATES)
    lc = [d for d in ev.diffs if d.kind == "lc"]
    assert len(lc) == len(sourced) * len(report.entities)


# --------------------------------------------------------------------------
# Requirement #2: TB net profit is the column-wise sum of the P&L-series GLs,
# computed from the accounts (not a fixed SUM row), excluding balance sheet.
# --------------------------------------------------------------------------

def _make_tb(path, extra_rows):
    """Clone the sample TB and append extra (account, chf, eur, dkk) rows."""
    wb = openpyxl.load_workbook(TB)
    ws = wb.active
    r = ws.max_row + 1
    for acct, chf, eur, dkk in extra_rows:
        ws.cell(r, 1, acct)            # Group Account Number (col A)
        ws.cell(r, 4, chf); ws.cell(r, 5, eur); ws.cell(r, 6, dkk)
        r += 1
    wb.save(path)


def test_tb_net_profit_excludes_balance_sheet(tmp_path):
    from plcheck.config import CheckConfig
    cfg = CheckConfig()
    accounts, _ = inputs.tb_values(TB, ["BALSCH", "BALSDE", "BALSDK"])
    expected = sum(v["BALSCH"] for a, v in accounts.items() if cfg.is_pl_account(a))

    # adding a balance-sheet account (>=400000) must NOT change the net profit
    tb2 = tmp_path / "TB_bs.xlsx"
    _make_tb(tb2, [(456000, 9_999_999, 0, 0)])
    acc2, _ = inputs.tb_values(str(tb2), ["BALSCH", "BALSDE", "BALSDK"])
    got = sum(v["BALSCH"] for a, v in acc2.items() if cfg.is_pl_account(a))
    assert round(got, 2) == round(expected, 2)


def test_tb_net_profit_includes_added_pl_gl(tmp_path):
    from plcheck.config import CheckConfig
    cfg = CheckConfig()
    accounts, _ = inputs.tb_values(TB, ["BALSCH", "BALSDE", "BALSDK"])
    base = sum(v["BALSCH"] for a, v in accounts.items() if cfg.is_pl_account(a))

    # a new P&L-series account (3-series) must flow into the net-profit sum
    tb2 = tmp_path / "TB_pl.xlsx"
    _make_tb(tb2, [(312000, 1000.0, 0, 0)])
    acc2, _ = inputs.tb_values(str(tb2), ["BALSCH", "BALSDE", "BALSDK"])
    got = sum(v["BALSCH"] for a, v in acc2.items() if cfg.is_pl_account(a))
    assert round(got - base, 2) == 1000.0


def test_is_pl_account_boundaries():
    from plcheck.config import CheckConfig
    cfg = CheckConfig()
    assert cfg.is_pl_account(110200) and cfg.is_pl_account(311025)
    assert not cfg.is_pl_account(445600)   # balance sheet
    assert not cfg.is_pl_account(885900)
    assert not cfg.is_pl_account(400000)   # 4-series excluded


def test_is_pl_account_text_formatted():
    """Account numbers stored as text (ERP exports) are still recognised."""
    from plcheck.config import CheckConfig
    cfg = CheckConfig()
    assert cfg.is_pl_account("110200")     # text P&L account
    assert cfg.is_pl_account(" 311025 ")   # text with stray spaces
    assert not cfg.is_pl_account("445600")
    assert not cfg.is_pl_account("885900")


def test_tb_values_handle_text_accounts(tmp_path):
    """TB net-profit sum is unaffected by accounts being stored as text."""
    import openpyxl as _xl
    num = inputs.tb_values(TB, ["BALSCH", "BALSDE", "BALSDK"])[0]

    # rewrite every account in the TB as a text string, then re-read
    wb = _xl.load_workbook(TB)
    ws = wb.active
    acct_col = inputs._find_cell(ws, "Group Account Number")[1]
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, acct_col)
        if isinstance(c.value, (int, float)):
            c.value = str(int(c.value))
    text_tb = tmp_path / "TB_text.xlsx"
    wb.save(text_tb)

    txt = inputs.tb_values(str(text_tb), ["BALSCH", "BALSDE", "BALSDK"])[0]
    cfg = __import__("plcheck.config", fromlist=["CheckConfig"]).CheckConfig()
    s_num = sum(v["BALSCH"] for a, v in num.items() if cfg.is_pl_account(a))
    s_txt = sum(v["BALSCH"] for a, v in txt.items() if cfg.is_pl_account(a))
    assert round(s_num, 2) == round(s_txt, 2)


# --------------------------------------------------------------------------
# Block-label robustness: GC - Total (and others) must still be checked when
# the export's spelling/spacing varies; if a block is truly absent, warn.
# --------------------------------------------------------------------------

def _clone_report_relabel(path, old_norm, new_text):
    """Clone the sample report, replacing a block label (row 2) by its
    normalized match with new_text."""
    from plcheck.config import normalize_label
    wb = openpyxl.load_workbook(REPORT)
    ws = wb.active
    for c in ws[2]:
        if c.value is not None and normalize_label(c.value) == old_norm:
            c.value = new_text
    wb.save(path)


def test_canonical_block_label_variants():
    from plcheck.config import canonical_block_label
    assert canonical_block_label("GC-Total") == "GC - Total"
    assert canonical_block_label("gc  -  total") == "GC - Total"
    assert canonical_block_label("LC - Balance ") == "LC - Balance"
    assert canonical_block_label("Mystery Block") == "Mystery Block"


def test_gc_total_check_runs_despite_spacing(tmp_path):
    from plcheck.workbook import build_check_workbook
    rep = tmp_path / "report_variant.xlsx"
    _clone_report_relabel(rep, "gc-total", "GC-Total")     # drop the spaces
    report = inputs.read_report(str(rep))
    assert report.block("GC - Total") is not None          # canonicalized
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)["Check"]
    assert str(ws["AJ9"].value).startswith("=SUMIF")       # GC-Total check built
    ev = evaluate(report, TB, AGG, RATES)
    assert "GC - Total" not in ev.missing_blocks


def test_missing_block_is_flagged(tmp_path):
    rep = tmp_path / "report_missing.xlsx"
    _clone_report_relabel(rep, "gc-total", "Grand Total XYZ")   # unrecognised
    report = inputs.read_report(str(rep))
    ev = evaluate(report, TB, AGG, RATES)                       # must not crash
    assert "GC - Total" in ev.missing_blocks


def test_rate_lookup_window_is_configurable():
    from plcheck.config import CheckConfig
    assert CheckConfig().rate_lookup_rows == 150


def _make_consol_tracker(path, rows):
    """Build a tracker in the real layout. rows = [(code, account, dr, cr)]."""
    import openpyxl as _xl
    wb = _xl.Workbook(); ws = wb.active; ws.title = "Entries- GR"
    ws.append(["Company", "Comp code", "Concatenate", "Group Account Number",
               "Functional Group", "GL Descriptions",
               "Feb v1", "Feb v2", "Mar v1", "Mar v2", "Currency", "Type"])
    for code, acct, dr, cr in rows:
        ws.append([code, code, f"{code}{acct}", acct, "COS", "entry",
                   "", "", dr, cr, "EUR", "BPC"])
    wb.save(path)


def test_consol_tracker_parsed_latest_month(tmp_path):
    p = tmp_path / "consol.xlsx"
    _make_consol_tracker(p, [("BALSCH", 290100, 222.0, ""),    # Dr leg (charge)
                             ("BALSCH", 804300, "", 222.0)])   # "To" Cr leg
    g = inputs.parse_consol(str(p))
    assert g.concat_col == 3
    assert len(g.val_cols) == 2                 # latest month's Dr/Cr pair
    # net posting = Debit - Credit: the Dr leg is +, the credit leg is -
    assert g.totals["BALSCH290100"] == 222.0
    assert g.totals["BALSCH804300"] == -222.0
    assert inputs.consol_key("BALSCH", 290100) == "BALSCH290100"


def test_entity_coverage_flags_source_only_code(tmp_path):
    """The Entity Coverage sheet lists codes per source and flags one present in
    the TB but missing from the report/agg (e.g. the ATI10N typo)."""
    out = tmp_path / "Check.xlsx"
    analyze(REPORT, TB, AGG, RATES, output_path=str(out))
    ws = openpyxl.load_workbook(out)["Entity Coverage"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row}
    flat = [v for v in cells.values() if isinstance(v, str)]
    assert "ATI10N" in flat                     # the TB-only code is listed
    # its reconciliation row marks it present only in the TB
    ati = next(r for (r, col), v in cells.items() if v == "ATI10N" and col >= 5)
    rowvals = [cells.get((ati, c)) for c in range(5, 10)]
    assert "ATI10N" in rowvals and rowvals.count("Yes") == 1 and "-" in rowvals


def test_company_codes_extracted(tmp_path):
    assert inputs.company_codes(TB) == ["ATI10N", "BALSCH", "BALSDE", "BALSDK"]
    assert "ATI10N" not in inputs.company_codes(AGG)


def test_dimension_label_not_a_company_code():
    """A dimension label like 'Consolidation unit' must not be taken as a code."""
    assert inputs.looks_like_company_code("BALSCH")
    assert inputs.looks_like_company_code("AT10IN")
    for label in ("Consolidation unit", "Company", "Overall Result",
                  "GL Description | Currency", "Group Account Number"):
        assert not inputs.looks_like_company_code(label)


def test_coverage_excludes_label_code(tmp_path):
    """A 'Consolidation unit' header sneaked into the TB is not listed."""
    import openpyxl as _xl
    p = tmp_path / "tb.xlsx"
    wb = _xl.Workbook(); ws = wb.active
    ws.cell(6, 2, "Company")
    for i, c in enumerate(["Consolidation unit", "BALSCH", "BALSDE"]):
        ws.cell(6, 3 + i, c)
    wb.save(p)
    assert inputs.company_codes(str(p)) == ["BALSCH", "BALSDE"]


def test_consol_functional_split(tmp_path):
    """A GL split across COS/S&M/G&A is summed per functional code (for a
    function-wise row) and as a single total (for a nature-wise row)."""
    import openpyxl as _xl
    from plcheck.config import CheckConfig
    p = tmp_path / "consol.xlsx"
    wb = _xl.Workbook(); ws = wb.active; ws.title = "Entries- GR"
    ws.append(["Company", "Comp code", "Concatenate", "Group Account Number",
               "Functional Group", "GL Descriptions",
               "Feb Dr", "Feb Cr", "Mar Dr", "Mar Cr", "Currency", "Type"])
    for fg, amt in [("COS", 100.0), ("S&M", 200.0), ("G&A", 300.0)]:
        ws.append(["BALSCH", "BALSCH", "BALSCH110200", 110200, fg, "x",
                   "", "", amt, "", "EUR", "BPC"])
    wb.save(p)
    g = inputs.parse_consol(str(p))
    assert g.func_col == 5
    # function-wise: each section ties to its own slice
    assert g.totals_fn[("BALSCH110200", "COS")] == 100.0
    assert g.totals_fn[("BALSCH110200", "S&M")] == 200.0
    assert g.totals_fn[("BALSCH110200", "G&A")] == 300.0
    # nature-wise: the single row = the total of all three
    assert g.totals["BALSCH110200"] == 600.0
    # category -> functional code mapping
    cfg = CheckConfig()
    assert cfg.functional_code("Cost of Production") == "COS"
    assert cfg.functional_code("Sales") == "S&M"
    assert cfg.functional_code("General Administration") == "G&A"
    assert cfg.functional_code("Revenue") is None        # not a functional split


def test_functional_code_inherited_from_contra_leg(tmp_path):
    """The COS/S&M/G&A code may be on the P&L leg OR the contra leg; the P&L
    leg must still match its functional slice."""
    import openpyxl as _xl
    p = tmp_path / "consol.xlsx"
    wb = _xl.Workbook(); ws = wb.active; ws.title = "Entries- GR"
    ws.append(["Company", "Comp code", "Concatenate", "Group Account Number",
               "Functional Group", "GL Descriptions",
               "Feb Dr", "Feb Cr", "Mar Dr", "Mar Cr", "Currency", "Type"])
    # P&L leg has a BLANK functional code; the contra leg carries "S&M"
    ws.append(["BALSCH", "BALSCH", "BALSCH110200", 110200, "", "charge",
               "", "", 1234.0, "", "CHF", "BPC"])
    ws.append(["BALSCH", "BALSCH", "BALSCH804340", 804340, "S&M", "To ...",
               "", "", "", 1234.0, "CHF", "BPC"])
    ws.append([None] * 12)
    wb.save(p)
    g = inputs.parse_consol(str(p))
    assert g.totals_fn.get(("BALSCH110200", "S&M")) == 1234.0


def test_consol_credit_leg_is_negative(tmp_path):
    """A 'To ...' credit-leg P&L account (amount in the Credit column) nets
    negative, matching the report's signed LC - Consol."""
    p = tmp_path / "consol.xlsx"
    _make_consol_tracker(p, [("BALSCH", 290900, "", 7000.0)])  # credit leg
    g = inputs.parse_consol(str(p))
    assert g.totals["BALSCH290900"] == -7000.0


def test_consol_tieout_flags_unbacked_entry(tmp_path):
    from plcheck.config import CheckConfig
    p = tmp_path / "consol.xlsx"
    # report LC-Consol for 290100/BALSCH is 0; tracker says 222 -> diff 222
    _make_consol_tracker(p, [("BALSCH", 290100, 222.0, "")])
    report = inputs.read_report(REPORT)
    ev = evaluate(report, TB, AGG, RATES, CheckConfig(), str(p))
    hits = [d for d in ev.flagged()
            if d.kind == "lc_consol" and d.account == 290100
            and d.entity == "BALSCH"]
    assert hits and abs(hits[0].expected - 222.0) < 1e-9
    assert hits[0].stated == 0.0


def test_consol_check_builds_formula_and_embeds_sheet(tmp_path):
    from plcheck.config import CheckConfig
    from plcheck.workbook import build_check_workbook
    p = tmp_path / "consol.xlsx"
    _make_consol_tracker(p, [("BALSCH", 290100, 222.0, "")])
    report = inputs.read_report(REPORT)
    # without a tracker: no LC-Consol diff column at all
    plain = tmp_path / "plain.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(plain)
    pw = openpyxl.load_workbook(plain)
    assert "Consol Entries" not in pw.sheetnames
    # with a tracker: tracker embedded + a dedicated LC-Consol Check sheet that
    # carries the SUMIF tie-out formulas (not the Check tab)
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES,
                         consol_path=str(p)).save(out)
    wb = openpyxl.load_workbook(out)
    assert "Consol Entries" in wb.sheetnames
    assert "LC-Consol Check" in wb.sheetnames
    ws = wb["LC-Consol Check"]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and "Consol Entries" in c.value
                and "SUMIF" in c.value]
    assert formulas and formulas[0].count("SUMIF") == 2   # the Dr + Cr columns
    # the Check tab no longer carries a consol diff
    assert not [c.value for row in wb["Check"].iter_rows() for c in row
                if isinstance(c.value, str) and "Consol Entries" in c.value]


def test_rates_columns_found_by_currency_data(tmp_path):
    """The 'From' currency column is found by its data, so a 'To Currency'
    column sitting to its left/right (all INR) can't be picked by mistake,
    regardless of column count or header wording."""
    import openpyxl as _xl
    p = tmp_path / "rates_real_layout.xlsx"
    wb = _xl.Workbook(); ws = wb.active
    # user's real layout: Type | From Currency | To Currency | Date | Exch Rate
    ws.append(["Exchange Rate Type", "From Currency", "To Currency",
               "Date", "Exchange Rate"])
    for cur, rt in [("AED", 25.51), ("CHF", 67.81), ("EUR", 107.82),
                    ("USD", 93.57)]:
        ws.append(["MA", cur, "INR", 31032026, rt])
    wb.save(p)
    g = inputs.parse_rates(str(p))
    assert g.from_col == 2          # 'From Currency' (B), not 'To Currency' (C)
    assert g.rate_col == 5          # 'Exchange Rate' (E), not 'Rate Type' (A)
    assert g.col_index == 4
    assert g.rates["CHF"] == 67.81 and "INR" not in g.rates


def test_base_currency_gc_rate_has_no_divisor(tmp_path):
    """When the group currency is the table's base (no row of its own), the
    GC-Balance rate is a plain lookup — not divided by a missing VLOOKUP."""
    from plcheck.workbook import build_check_workbook
    import openpyxl as _xl
    # rates with NO 'INR' row (INR is only the To-currency / base)
    rp = tmp_path / "rates_no_inr.xlsx"
    wb = _xl.Workbook(); ws = wb.active
    ws.append(["Exchange Rate Type", "From Currency", "To Currency",
               "Date", "Exchange Rate"])
    for cur, rt in [("CHF", 67.81), ("EUR", 107.82), ("DKK", 14.22)]:
        ws.append(["MA", cur, "INR", 31032026, rt])
    wb.save(rp)
    report = inputs.read_report(REPORT)
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, str(rp)).save(out)
    cell = openpyxl.load_workbook(out)["Check"]["Q3"].value
    assert cell.startswith("=VLOOKUP(") and "/VLOOKUP(" not in cell


def test_minority_interest_matched_fuzzily():
    """A misspelled Minority-Interest section is still recognised + excluded."""
    from plcheck.config import CheckConfig, CategoryRule
    cfg = CheckConfig()
    for name in ("Minority Interest", "Minority interest", "Mintority Interest",
                 "Minority  Interest", "MINORITY INTERST",
                 "Non-controlling Interest"):
        assert cfg.is_minority(name), name
        # excluded from the P&L (blank classification), never None/unmapped
        assert cfg.effective_rule(name) == CategoryRule("", "")
    # unrelated sections must NOT be mistaken for Minority Interest
    for name in ("Interest", "Finance Cost", "Other Income", "Revenue"):
        assert not cfg.is_minority(name), name


def test_minority_typo_not_swept_into_nature_wise_expense():
    """In Nature-wise mode the default expense rule must not swallow a typo'd
    Minority Interest — it stays excluded from the P&L."""
    from plcheck.config import CheckConfig, CategoryRule
    cfg = CheckConfig()
    cfg.default_rule = CategoryRule("Expense", "tb")     # Nature-wise fallback
    assert cfg.effective_rule("Mintority Interest") == CategoryRule("", "")
    assert cfg.effective_rule("Travel expenses") == CategoryRule("Expense", "tb")


# --------------------------------------------------------------------------
# Ind-AS Function-wise report: different section names + a Depreciation section
# (which may contain multiple GLs).
# --------------------------------------------------------------------------

INDAS = os.path.join(SAMPLE, "INDAS_Functionwise_PL_Report.xlsx")


def test_indas_report_reconciles():
    report = inputs.read_report(INDAS)
    ev = evaluate(report, TB, AGG, RATES)
    assert ev.unmapped_categories == []          # all INDAS sections mapped
    assert ev.missing_blocks == []               # all blocks recognised
    # internal P&L integrity holds for every entity
    assert all(abs(n.calc_check) < 0.01 for n in ev.net_profit)
    # the only LC tie-out breaks are the genuine depreciation roundings
    lc_bad = {(d.category, d.account) for d in ev.flagged() if d.kind == "lc"}
    assert lc_bad == {("Depreciation", 290100)}


def test_depreciation_section_supports_multiple_gls(tmp_path):
    """A second GL in the Depreciation section is reconciled like the first."""
    from plcheck.model import ReportRow
    from plcheck.workbook import build_check_workbook, FIRST_DATA_ROW

    report = inputs.read_report(INDAS)
    new_row = ReportRow(
        category="Depreciation", account=290200,
        description="Depreciation - Plant & Machinery",
        values={(b.label, sub): 111.0
                for b in report.blocks for sub in b.columns},
    )
    pos = max(i for i, r in enumerate(report.rows)
              if r.category == "Depreciation") + 1
    report.rows.insert(pos, new_row)

    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)["Check"]
    r = FIRST_DATA_ROW + pos
    assert ws[f"C{r}"].value == 290200
    # ties back to the Real Time TB, by account number, like the first dep. GL
    assert str(ws[f"F{r}"].value).startswith("=IFERROR(VLOOKUP")
    assert "'Real Time TB'" in ws[f"F{r}"].value
    assert str(ws[f"Q{r}"].value).startswith("=SUMIF")    # FX
    assert str(ws[f"AJ{r}"].value).startswith("=SUMIF")   # consolidation


def test_both_report_styles_share_one_config():
    """IFRS and INDAS section names both resolve with the default config."""
    from plcheck.config import CheckConfig
    cfg = CheckConfig()
    for cat in ("Revenue", "Cost of Production", "Sales", "General Administration"):
        assert cfg.rule_for(cat) is not None
    for cat in ("Income", "Software Development Exp", "Sales & Marketing Cost",
                "Administration cost", "Depreciation"):
        assert cfg.rule_for(cat) is not None


# --------------------------------------------------------------------------
# An entity in the report that has no matching company column in the TB must
# get 0 for "Net Profit as per Real Time TB" (not the first company's figure),
# and be flagged.
# --------------------------------------------------------------------------

def _clone_report_rename_entity(path, old_code, new_code):
    """Clone the sample report, renaming an entity code across the sub-header
    row (row 3) so it no longer matches the TB."""
    wb = openpyxl.load_workbook(REPORT)
    ws = wb.active
    for c in ws[3]:
        if isinstance(c.value, str) and c.value.strip() == old_code:
            c.value = new_code
    wb.save(path)


def test_unmatched_entity_tb_net_profit_is_zero(tmp_path):
    from plcheck.workbook import build_check_workbook
    rep = tmp_path / "report_badcode.xlsx"
    _clone_report_rename_entity(rep, "BALSDK", "BALSXX")   # not in the TB
    report = inputs.read_report(str(rep))

    # flagged in the summary
    ev = evaluate(report, TB, AGG, RATES)
    assert "BALSXX" in ev.missing_entities

    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)["Check"]
    # locate the BALSXX value column and its diff (output) column
    from plcheck.layout import plan
    from plcheck import config as C
    layout = plan(report)
    d = layout.diff_col(C.CHECK_LC_BALANCE, "BALSXX")
    # find the "Net Profit as per Real Time TB" row
    row = next(r for r in range(33, 50)
               if ws[f"D{r}"].value == "Net Profit as per Real Time TB")
    assert ws[f"{d}{row}"].value == 0          # not a SUMPRODUCT, not AUD's value

    # a matched entity still gets the live SUMPRODUCT
    d_ok = layout.diff_col(C.CHECK_LC_BALANCE, "BALSCH")
    assert str(ws[f"{d_ok}{row}"].value).startswith("=SUMPRODUCT")


def test_entity_matching_is_case_insensitive(tmp_path):
    rep = tmp_path / "report_lower.xlsx"
    _clone_report_rename_entity(rep, "BALSDK", "balsdk")   # same code, lowercase
    report = inputs.read_report(str(rep))
    ev = evaluate(report, TB, AGG, RATES)
    assert ev.missing_entities == []           # still matched to the TB


# --------------------------------------------------------------------------
# Minority Interest sheet derived from the Check sheet.
# --------------------------------------------------------------------------

def test_minority_interest_sheet(tmp_path):
    from plcheck.workbook import build_check_workbook, MINORITY_SHEET
    report = inputs.read_report(REPORT)
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    wb = openpyxl.load_workbook(out)
    assert MINORITY_SHEET in wb.sheetnames
    ws = wb[MINORITY_SHEET]
    # one row per company code from the LC section
    codes = [ws.cell(2 + k, 1).value for k in range(len(report.entities))]
    assert codes == [e.code for e in report.entities]
    assert ws.cell(1, 1).value == "Code"
    assert ws.cell(1, 6).value == "Current Period %"
    # B = GC-Balance Net Profit, E = GC-Total Minority, D = B+C, F = E/D
    assert ws.cell(2, 2).value.startswith("='Check'!")
    assert ws.cell(2, 4).value == "=B2+C2"
    assert ws.cell(2, 5).value.startswith("='Check'!")
    assert ws.cell(2, 6).value == "=IFERROR(E2/D2,0)"


def test_minority_dividend_account_picked_up(tmp_path):
    from plcheck.model import ReportRow
    from plcheck.workbook import build_check_workbook, MINORITY_SHEET
    report = inputs.read_report(REPORT)
    # a dividend-received GL 332010 should feed column C from GC-Balance
    report.rows.insert(0, ReportRow(
        category="Other Income", account=332010, description="Dividend received",
        values={(b.label, sub): 0.0
                for b in report.blocks for sub in b.columns}))
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)[MINORITY_SHEET]
    assert str(ws.cell(2, 3).value).startswith("='Check'!")   # not the literal 0


def test_minority_interest_for_indas(tmp_path):
    from plcheck.workbook import build_check_workbook, MINORITY_SHEET
    report = inputs.read_report(INDAS)
    out = tmp_path / "Check.xlsx"
    build_check_workbook(report, TB, AGG, RATES).save(out)
    assert MINORITY_SHEET in openpyxl.load_workbook(out).sheetnames


# --------------------------------------------------------------------------
# Header band that starts on row 1 (some exports) instead of row 2 must be
# detected, not assumed.
# --------------------------------------------------------------------------

ROW1 = os.path.join(SAMPLE, "IFRS_INR_PL_Report_row1header.xlsx")


def test_report_header_on_row1_is_detected():
    report = inputs.read_report(ROW1)
    # entity codes (not the company names) must be read correctly
    assert [e.code for e in report.entities] == ["BALSCH", "BALSDE", "BALSDK"]
    assert [e.currency for e in report.entities] == ["CHF", "EUR", "DKK"]


def test_row1header_report_reconciles(tmp_path):
    report = inputs.read_report(ROW1)
    ev = evaluate(report, TB, AGG, RATES)
    assert ev.unmapped_categories == [] and ev.missing_blocks == []
    # net profit ties except the genuine depreciation rounding
    tie = {n.entity: round(n.tie_check, 2) for n in ev.net_profit}
    assert tie["BALSCH"] == 0.0 and tie["BALSDE"] == 0.51 and tie["BALSDK"] == 0.65
    out = tmp_path / "Check.xlsx"
    analyze(ROW1, TB, AGG, RATES, output_path=str(out))
    assert out.is_file()


# --------------------------------------------------------------------------
# A report cell holding an Excel error (e.g. #REF! from a broken formula) must
# be surfaced, not silently treated as 0.
# --------------------------------------------------------------------------

def test_report_excel_error_cell_is_flagged(tmp_path):
    wb = openpyxl.load_workbook(REPORT)
    ws = wb.active
    ws["AB6"] = "#REF!"          # a GC-Total cell on a data row
    p = tmp_path / "err.xlsx"
    wb.save(p)
    report = inputs.read_report(str(p))
    assert any("#REF!" in e for e in report.errors)
    ev = evaluate(report, TB, AGG, RATES)
    assert any("#REF!" in e for e in ev.report_errors)


# --------------------------------------------------------------------------
# Ind-AS Nature-wise: no Aggregate Exp; every line ties to the Real Time TB.
# --------------------------------------------------------------------------

NATUREWISE = os.path.join(SAMPLE, "INDAS_Naturewise_PL_Report.xlsx")


def _naturewise_cfg():
    from plcheck.config import CheckConfig, CategoryRule
    return CheckConfig(default_rule=CategoryRule("Expense", "tb"))


def test_naturewise_runs_without_agg(tmp_path):
    cfg = _naturewise_cfg()
    report = inputs.read_report(NATUREWISE)
    ev = evaluate(report, TB, None, RATES, cfg)      # no Aggregate Exp
    assert ev.unmapped_categories == []
    assert all(abs(n.calc_check) < 0.01 for n in ev.net_profit)
    tie = {n.entity: round(n.tie_check, 2) for n in ev.net_profit}
    assert tie["BALSCH"] == 0.0 and tie["BALSDE"] == 0.51 and tie["BALSDK"] == 0.65


def test_naturewise_output_has_no_aggregate_sheet(tmp_path):
    cfg = _naturewise_cfg()
    out = tmp_path / "Check.xlsx"
    analyze(NATUREWISE, TB, None, RATES, output_path=str(out), cfg=cfg)
    sheets = openpyxl.load_workbook(out).sheetnames
    assert "Aggregate Exp" not in sheets
    assert sheets[0] == "Check"
    ws = openpyxl.load_workbook(out)["Check"]
    # an expense nature (Employee Benefit Expenses, GL 110200) ties to the TB
    assert "'Real Time TB'" in str(ws["F10"].value)


def test_default_rule_classifies_unmapped_as_expense():
    from plcheck.config import CategoryRule
    cfg = _naturewise_cfg()
    assert cfg.effective_rule("Some New Expense Line") == CategoryRule("Expense", "tb")
    assert cfg.effective_rule("Income").classification == "Income"   # explicit wins


# --------------------------------------------------------------------------
# IFRS USD: group currency is USD, so GC = local->INR / USD->INR (cross-rate),
# auto-detected from the report.
# --------------------------------------------------------------------------

USD_REPORT = os.path.join(SAMPLE, "IFRS_USD_PL_Report.xlsx")


def test_gc_currency_detection():
    rates = inputs.parse_rates(RATES).rates
    assert inputs.detect_gc_currency(inputs.read_report(REPORT), rates) == "INR"
    assert inputs.detect_gc_currency(inputs.read_report(USD_REPORT), rates) == "USD"


def test_usd_report_reconciles(tmp_path):
    rep = inputs.read_report(USD_REPORT)
    ev = evaluate(rep, TB, AGG, RATES)
    # net profit (LC terms) ties internally, same as the INR report
    assert all(abs(n.calc_check) < 0.01 for n in ev.net_profit)
    # the depreciation rounding now shows in USD (~0.59), not INR (~54.99)
    fx = {round(d.delta, 2) for d in ev.flagged() if d.kind == "fx"}
    assert 0.59 in fx and 54.99 not in fx
    out = tmp_path / "Check.xlsx"
    analyze(USD_REPORT, TB, AGG, RATES, output_path=str(out))
    assert out.is_file()


def test_usd_rate_formula_uses_cross_rate(tmp_path):
    from plcheck.workbook import build_check_workbook
    out = tmp_path / "Check.xlsx"
    build_check_workbook(inputs.read_report(USD_REPORT), TB, AGG, RATES).save(out)
    ws = openpyxl.load_workbook(out)["Check"]
    # the GC-Balance rate helper divides by the USD->INR rate
    assert 'VLOOKUP("USD"' in str(ws["Q3"].value)
