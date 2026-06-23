"""Tests for the PDF → Excel converter.

These build a tiny statement PDF on the fly (reportlab) and round-trip it
through the converter, so they only run where the optional PDF/Excel libraries
are installed.
"""

import pytest

pdfplumber = pytest.importorskip("pdfplumber")
openpyxl = pytest.importorskip("openpyxl")
reportlab = pytest.importorskip("reportlab")

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from fincheck.to_excel import convert_to_excel, extract_grids


def _make_pdf(path):
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    styles = getSampleStyleSheet()
    data = [
        ["Particulars", "31 Mar 2024", "31 Mar 2023"],
        ["Revenue from operations", "1,24,500", "1,08,300"],
        ["Other income", "3,450", "2,100"],
        ["Total income", "1,27,950", "1,10,400"],
        ["Finance costs", "(2,300)", "(1,950)"],
        ["Profit before tax", "INR 41,250", "INR 33,150"],
        ["Nil item", "-", "-"],
    ]
    table = Table(data, colWidths=[8 * cm, 3 * cm, 3 * cm])
    table.setStyle(
        TableStyle(
            [
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
            ]
        )
    )
    story = [
        Paragraph("XYZ Ltd — Statement of Profit and Loss", styles["Heading2"]),
        Spacer(1, 0.4 * cm),
        table,
    ]
    doc.build(story)


def _cells_by_label(ws):
    """Map the first-column label of each row to the row's remaining cells."""
    out = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0]:
            out[row[0]] = row[1:]
    return out


def test_converts_and_parses_numbers(tmp_path):
    pdf = tmp_path / "stmt.pdf"
    _make_pdf(pdf)
    out = convert_to_excel(str(pdf), str(tmp_path / "stmt.xlsx"))

    wb = load_workbook(out)
    assert "All Data" in wb.sheetnames
    assert "Page 1" in wb.sheetnames

    rows = _cells_by_label(wb["Page 1"])
    # Indian-grouped thousands become real numbers.
    assert rows["Revenue from operations"][:2] == (124500, 108300)
    # Parenthesised negatives.
    assert rows["Finance costs"][:2] == (-2300, -1950)
    # Currency-coded figures are still parsed.
    assert rows["Profit before tax"][:2] == (41250, 33150)


def test_dates_kept_as_text(tmp_path):
    pdf = tmp_path / "stmt.pdf"
    _make_pdf(pdf)
    out = convert_to_excel(str(pdf), str(tmp_path / "stmt.xlsx"))
    rows = _cells_by_label(load_workbook(out)["Page 1"])
    assert rows["Particulars"][0] == "31 Mar 2024"


def test_raw_mode_keeps_original_text(tmp_path):
    pdf = tmp_path / "stmt.pdf"
    _make_pdf(pdf)
    out = convert_to_excel(str(pdf), str(tmp_path / "raw.xlsx"), raw=True)
    rows = _cells_by_label(load_workbook(out)["Page 1"])
    assert rows["Revenue from operations"][0] == "1,24,500"
    assert rows["Finance costs"][0] == "(2,300)"


def test_sheet_toggles(tmp_path):
    pdf = tmp_path / "stmt.pdf"
    _make_pdf(pdf)
    out = convert_to_excel(
        str(pdf), str(tmp_path / "p.xlsx"), combined=False, per_page=True
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["Page 1"]


def test_extract_grids_structure(tmp_path):
    pdf = tmp_path / "stmt.pdf"
    _make_pdf(pdf)
    grids = extract_grids(str(pdf))
    assert len(grids) == 1
    assert grids[0].n_columns == 3
    labels = [r[0] for r in grids[0].rows]
    assert "Total income" in labels
