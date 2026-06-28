"""Convert an Excel workbook into CSV files — one CSV per worksheet.

A workbook (``.xlsx`` / ``.xlsm``) can hold many sheets, so "converting it to
CSV" really means writing *a folder of CSVs*, one per sheet. This module does
that robustly for real-world financial workbooks:

* every sheet becomes ``<name>.csv`` (sheet names are slugified so they are
  safe filenames), or you can pick a single sheet;
* formula cells are written as their **last cached value** (what Excel showed),
  not the ``=SUM(...)`` text;
* dates/times are emitted in ISO format instead of Excel serial numbers;
* merged cells repeat the top-left value across the merged range so columns
  still line up (xlsx only — merge info is unavailable in streaming readers);
* trailing empty rows and columns are trimmed;
* files are written UTF-8 with a BOM by default so Excel re-opens them cleanly,
  and the delimiter / encoding / line ending are all configurable.

Public API::

    from fincheck.xlsx2csv import convert_workbook
    written = convert_workbook("book.xlsx", "out_dir")   # -> list[Path]

Command line::

    python -m fincheck.xlsx2csv book.xlsx -o out_dir
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import re
import sys
from pathlib import Path

try:  # openpyxl is the only hard dependency of this module.
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:  # pragma: no cover - import guard
    raise ModuleNotFoundError(
        "Converting Excel workbooks needs 'openpyxl'. Install it with "
        "`pip install openpyxl` (or `pip install -r requirements.txt`)."
    ) from exc


_SLUG_RE = re.compile(r"[^A-Za-z0-9._-]+")


def _slugify(name: str, fallback: str) -> str:
    """Turn a sheet name into a safe, non-empty filename stem."""
    slug = _SLUG_RE.sub("_", name).strip("._-")
    return slug or fallback


def _format_cell(value):
    """Render a single cell value the way a person would expect in a CSV."""
    if value is None:
        return ""
    if isinstance(value, bool):  # bool is an int subclass — check it first.
        return "TRUE" if value else "FALSE"
    if isinstance(value, _dt.datetime):
        # Drop a midnight time component so pure dates stay dates.
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, _dt.time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        # 1234.0 -> "1234" so integer-valued floats don't grow a ".0".
        return str(int(value))
    return str(value)


def _sheet_rows(ws):
    """Yield trimmed rows of formatted strings for a worksheet.

    Merged cells contribute their value only in the top-left position in
    openpyxl; we copy it across the range so the data stays aligned.
    """
    merged_value: dict[tuple[int, int], object] = {}
    # merged_cells is unavailable on read-only worksheets; degrade gracefully.
    for rng in getattr(getattr(ws, "merged_cells", None), "ranges", ()):
        top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
        if top_left is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_value[(r, c)] = top_left

    rows: list[list[str]] = []
    max_width = 0
    for r, row in enumerate(ws.iter_rows(), start=1):
        out: list[str] = []
        for c, cell in enumerate(row, start=1):
            value = merged_value.get((r, c), cell.value)
            out.append(_format_cell(value))
        # Trim trailing empty cells on this row.
        while out and out[-1] == "":
            out.pop()
        max_width = max(max_width, len(out))
        rows.append(out)

    # Drop trailing fully-empty rows.
    while rows and not any(rows[-1]):
        rows.pop()

    # Pad every row to the widest one so the CSV is rectangular.
    for row in rows:
        if len(row) < max_width:
            row.extend([""] * (max_width - len(row)))
    return rows


def convert_workbook(
    source,
    output_dir=None,
    *,
    sheet=None,
    delimiter=",",
    encoding="utf-8-sig",
    line_terminator="\r\n",
    skip_empty_sheets=True,
    overwrite=True,
):
    """Convert an Excel workbook to one CSV per worksheet.

    Args:
        source: path to the ``.xlsx`` / ``.xlsm`` workbook.
        output_dir: folder to write CSVs into (default: ``<source stem>_csv``
            next to the workbook). Created if it does not exist.
        sheet: only convert this sheet (name) instead of all of them.
        delimiter: field separator for the CSV (default ``,``).
        encoding: output encoding (default ``utf-8-sig`` so Excel opens it).
        line_terminator: row separator written into the file.
        skip_empty_sheets: do not emit a CSV for a sheet that has no data.
        overwrite: overwrite existing CSVs; if ``False``, raise on collision.

    Returns:
        list[Path]: the CSV files that were written.
    """
    src = Path(source)
    if not src.is_file():
        raise FileNotFoundError(f"workbook not found: {src}")

    out_dir = Path(output_dir) if output_dir is not None else src.with_name(
        f"{src.stem}_csv"
    )
    out_dir.mkdir(parents=True, exist_ok=True)

    # data_only -> write cached formula results, not the "=SUM(...)" text.
    wb = load_workbook(src, data_only=True)
    try:
        names = wb.sheetnames
        if sheet is not None:
            if sheet not in names:
                raise KeyError(
                    f"sheet {sheet!r} not in workbook; available: {names}"
                )
            names = [sheet]

        written: list[Path] = []
        used_stems: dict[str, int] = {}
        for index, name in enumerate(names):
            ws = wb[name]
            rows = _sheet_rows(ws)
            if skip_empty_sheets and not rows:
                continue

            # Build a unique, safe filename stem for this sheet.
            stem = _slugify(name, fallback=f"sheet{index + 1}")
            count = used_stems.get(stem.lower(), 0)
            used_stems[stem.lower()] = count + 1
            if count:
                stem = f"{stem}_{count + 1}"

            target = out_dir / f"{stem}.csv"
            if target.exists() and not overwrite:
                raise FileExistsError(f"refusing to overwrite {target}")

            with open(target, "w", newline="", encoding=encoding) as fh:
                writer = csv.writer(
                    fh, delimiter=delimiter, lineterminator=line_terminator
                )
                writer.writerows(rows)
            written.append(target)
        return written
    finally:
        wb.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="fincheck-xlsx2csv",
        description="Convert an Excel workbook into CSV files (one per sheet).",
    )
    parser.add_argument("workbook", help="path to the .xlsx / .xlsm workbook")
    parser.add_argument(
        "-o",
        "--output",
        help="output directory for the CSVs "
        "(default: <workbook>_csv next to the file)",
    )
    parser.add_argument(
        "-s",
        "--sheet",
        help="convert only this sheet (by name) instead of all of them",
    )
    parser.add_argument(
        "-d",
        "--delimiter",
        default=",",
        help=r"field delimiter (default ','; use $'\t' for TSV)",
    )
    parser.add_argument(
        "--encoding",
        default="utf-8-sig",
        help="output encoding (default utf-8-sig, which Excel opens cleanly)",
    )
    parser.add_argument(
        "--keep-empty",
        action="store_true",
        help="also write a CSV for sheets that contain no data",
    )
    parser.add_argument(
        "--no-clobber",
        action="store_true",
        help="error instead of overwriting an existing CSV",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        written = convert_workbook(
            args.workbook,
            output_dir=args.output,
            sheet=args.sheet,
            delimiter=args.delimiter,
            encoding=args.encoding,
            skip_empty_sheets=not args.keep_empty,
            overwrite=not args.no_clobber,
        )
    except (FileNotFoundError, KeyError, FileExistsError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    if not written:
        print("no sheets with data found — nothing written.", file=sys.stderr)
        return 1

    print(f"Wrote {len(written)} CSV file(s):")
    for path in written:
        print(f"  {path}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
