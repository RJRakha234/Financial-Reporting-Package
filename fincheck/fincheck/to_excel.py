"""Convert a financial-statement / XBRL-filing PDF into an Excel workbook.

Many XBRL filings are circulated as a *rendered PDF* (e.g. an MCA AOC-4 / IFRS
filing printed to PDF). Validating the numbers in that form is painful, so this
module reconstructs the tabular grid of every page and writes it to ``.xlsx`` —
with figures stored as **real Excel numbers** (parentheses → negatives,
thousands separators stripped, currency symbols/nil dashes handled) so you can
sum, cross-foot and reconcile the data directly in Excel.

The grid is recovered without relying on ruling lines: words are clustered into
rows by vertical position, and column boundaries are found from the vertical
*whitespace gaps* that persist across the rows of a page (the way the eye reads a
column). This keeps both label text *and* text-valued fields (dates, yes/no,
names) in their own cells, not just the numeric columns.

Public API::

    from fincheck import convert_to_excel
    convert_to_excel("filing.pdf", "filing.xlsx")
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber

from .extract import _ROW_TOLERANCE, _cluster_rows, _merge_numberish
from .numbers import parse_number

# A vertical whitespace band at least this wide (points) can separate columns.
_MIN_GAP = 5.0
# A bin counts as part of a column separator only if it is whitespace in at
# least this fraction of the page's data rows. Tolerating a small fraction lets
# the occasional full-width heading or merged cell cross a separator without
# destroying the column it sits in.
_SEPARATOR_WHITESPACE_FRACTION = 0.90
# Labels that mark a row as a total/subtotal worth bolding in the output.
_TOTAL_WORDS = ("total", "subtotal", "sub-total", "aggregate", "grand total")


@dataclass
class PageGrid:
    """A single PDF page reconstructed as a rectangular grid of text cells."""

    index: int
    rows: list[list[str]] = field(default_factory=list)

    @property
    def n_columns(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def _content_bounds(words: list[dict]) -> tuple[float, float]:
    return min(w["x0"] for w in words), max(w["x1"] for w in words)


def _column_boundaries(
    row_words: list[list[dict]], x0: float, x1: float, n_rows: int
) -> list[float]:
    """Find column-separator x-positions from persistent vertical whitespace.

    A 1-point-resolution occupancy histogram is built across the page's content
    width: ``blocked[i]`` counts how many rows have ink over bin ``i``. Runs of
    bins that are whitespace in nearly every row, and at least ``_MIN_GAP`` wide,
    are the gutters between columns; each gutter's midpoint is a boundary.
    """
    span = x1 - x0
    if span <= 0:
        return [x0, x1]
    n_bins = int(math.ceil(span)) + 1
    blocked = [0] * n_bins
    for words in row_words:
        covered = bytearray(n_bins)
        for w in words:
            a = max(0, int(w["x0"] - x0))
            b = min(n_bins - 1, int(math.ceil(w["x1"] - x0)))
            for i in range(a, b + 1):
                covered[i] = 1
        for i in range(n_bins):
            blocked[i] += covered[i]

    # Max rows that may have ink in a bin for it to still count as whitespace.
    allowed = (1.0 - _SEPARATOR_WHITESPACE_FRACTION) * n_rows
    is_gap = [blocked[i] <= allowed for i in range(n_bins)]

    boundaries = [x0]
    run_start = None
    for i in range(n_bins):
        if is_gap[i]:
            if run_start is None:
                run_start = i
        else:
            if run_start is not None and i - run_start >= _MIN_GAP:
                boundaries.append(x0 + (run_start + i) / 2.0)
            run_start = None
    boundaries.append(x1 + 1.0)
    return boundaries


def _assign_to_columns(words: list[dict], boundaries: list[float]) -> list[str]:
    """Slot each word into the column whose band contains its centre."""
    n_cols = len(boundaries) - 1
    buckets: list[list[dict]] = [[] for _ in range(n_cols)]
    for w in words:
        centre = (w["x0"] + w["x1"]) / 2.0
        col = 0
        for c in range(n_cols):
            if boundaries[c] <= centre < boundaries[c + 1]:
                col = c
                break
        else:
            col = n_cols - 1
        buckets[col].append(w)
    cells = []
    for bucket in buckets:
        bucket.sort(key=lambda w: w["x0"])
        cells.append(" ".join(w["text"] for w in bucket).strip())
    return cells


def _build_grid(page_index: int, page) -> PageGrid:
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
    if not words:
        return PageGrid(index=page_index, rows=[])

    clustered = _cluster_rows(words)
    clustered.sort(key=lambda row: min(w["top"] for w in row))
    row_words = [_merge_numberish(row) for row in clustered]

    x0, x1 = _content_bounds(words)
    boundaries = _column_boundaries(row_words, x0, x1, len(row_words))

    rows = [_assign_to_columns(row, boundaries) for row in row_words]
    # Drop columns that no row actually uses. A lone wide heading whose edge
    # sticks out past the table body can otherwise leave an all-blank column.
    width = max((len(r) for r in rows), default=0)
    used = [c for c in range(width) if any(c < len(r) and r[c] for r in rows)]
    rows = [[r[c] if c < len(r) else "" for c in used] for r in rows]
    return PageGrid(index=page_index, rows=[r for r in rows if any(r)])


def extract_grids(pdf_path: str) -> list[PageGrid]:
    """Reconstruct every page of *pdf_path* as a :class:`PageGrid`."""
    grids: list[PageGrid] = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages):
            grids.append(_build_grid(i, page))
    return grids


def _as_cell_value(text: str):
    """Return a float for genuine figures, otherwise the original text.

    Only digit-bearing cells are even considered, so nil dashes (``-``/``nil``)
    and label text stay verbatim, while figures — including currency-coded ones
    like ``INR 41,250`` and parenthesised negatives ``(1,234)`` — become
    arithmetic-ready numbers. ``parse_number`` returns ``None`` for anything
    that is not ultimately a clean number (e.g. ``Note 21``, ``31 Mar 2024``),
    so those are preserved as text.
    """
    if any(ch.isdigit() for ch in text):
        value = parse_number(text)
        if value is not None:
            return value
    return text


def _is_total_row(cells: list[str]) -> bool:
    label = next((c for c in cells if c.strip()), "").lower()
    return any(word in label for word in _TOTAL_WORDS)


def _write_sheet(ws, rows: list[list[str]], raw: bool, header: list[str] | None):
    from openpyxl.styles import Font

    bold = Font(bold=True)
    max_widths: dict[int, int] = {}

    start = 1
    if header:
        for c, title in enumerate(header, 1):
            cell = ws.cell(row=1, column=c, value=title)
            cell.font = bold
            max_widths[c] = len(str(title))
        ws.freeze_panes = "A2"
        start = 2

    for r, cells in enumerate(rows, start):
        total_row = _is_total_row(cells)
        for c, text in enumerate(cells, 1):
            if not text:
                continue
            value = text if raw else _as_cell_value(text)
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, float):
                cell.number_format = (
                    "#,##0" if value == int(value) else "#,##0.00"
                )
            if total_row:
                cell.font = bold
            max_widths[c] = max(max_widths.get(c, 0), len(str(text)))

    for c, w in max_widths.items():
        from openpyxl.utils import get_column_letter

        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 70)


def grids_to_workbook(grids: list[PageGrid], raw: bool = False,
                      combined: bool = True, per_page: bool = True):
    """Build an openpyxl workbook from extracted page grids."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    if combined:
        ws = wb.create_sheet("All Data")
        width = max((g.n_columns for g in grids), default=0)
        header = ["Page"] + [f"Col {i + 1}" for i in range(width)]
        merged: list[list[str]] = []
        for g in grids:
            for row in g.rows:
                merged.append([str(g.index + 1)] + row)
        _write_sheet(ws, merged, raw, header)

    if per_page:
        for g in grids:
            if not g.rows:
                continue
            ws = wb.create_sheet(f"Page {g.index + 1}")
            _write_sheet(ws, g.rows, raw, header=None)

    if not wb.sheetnames:  # nothing extracted — keep a valid, non-empty file
        wb.create_sheet("All Data")
    return wb


def convert_to_excel(
    pdf_path: str,
    xlsx_path: str | None = None,
    *,
    raw: bool = False,
    combined: bool = True,
    per_page: bool = True,
) -> str:
    """Convert a PDF to an Excel workbook and return the output path.

    Args:
        pdf_path: the source PDF (a rendered XBRL filing / financial statement).
        xlsx_path: output path; defaults to ``<input>.xlsx``.
        raw: keep every value as the original text instead of converting
            figures to real numbers.
        combined: include an "All Data" sheet with every row across all pages.
        per_page: include one sheet per PDF page.
    """
    out = xlsx_path or str(Path(pdf_path).with_suffix(".xlsx"))
    grids = extract_grids(pdf_path)
    wb = grids_to_workbook(grids, raw=raw, combined=combined, per_page=per_page)
    wb.save(out)
    return out


def build_parser():
    import argparse

    parser = argparse.ArgumentParser(
        prog="fincheck.to_excel",
        description="Convert a financial-statement / XBRL-filing PDF into an "
        "Excel workbook with figures stored as real numbers for validation.",
    )
    parser.add_argument("pdf", help="path to the PDF to convert")
    parser.add_argument(
        "-o", "--output",
        help="output .xlsx path (default: <input>.xlsx)",
    )
    parser.add_argument(
        "--raw", action="store_true",
        help="keep values as original text instead of converting figures to "
        "numbers",
    )
    parser.add_argument(
        "--no-combined", action="store_true",
        help="omit the single 'All Data' sheet spanning every page",
    )
    parser.add_argument(
        "--no-pages", action="store_true",
        help="omit the per-page sheets",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    import sys

    args = build_parser().parse_args(argv)
    src = Path(args.pdf)
    if not src.is_file():
        print(f"error: file not found: {src}", file=sys.stderr)
        return 2
    if args.no_combined and args.no_pages:
        print("error: --no-combined and --no-pages leave nothing to write",
              file=sys.stderr)
        return 2

    out = args.output or str(src.with_suffix(".xlsx"))
    grids = extract_grids(str(src))
    wb = grids_to_workbook(
        grids,
        raw=args.raw,
        combined=not args.no_combined,
        per_page=not args.no_pages,
    )
    wb.save(out)

    pages = sum(1 for g in grids if g.rows)
    rows = sum(len(g.rows) for g in grids)
    print(f"Converted {pages} page(s), {rows} row(s) → {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
